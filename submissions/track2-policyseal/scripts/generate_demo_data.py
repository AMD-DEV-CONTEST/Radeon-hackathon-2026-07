from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import tempfile
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path

import fitz
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXED_TIME = datetime(2026, 1, 1, tzinfo=UTC)
ZIP_TIME = (1980, 1, 1, 0, 0, 0)
CORE_TIMESTAMP = b"2026-01-01T00:00:00Z"
CORE_TIMESTAMP_PATTERN = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+(</dcterms:(?:created|modified)>)"
)
PDF_TIMESTAMP = b"D:20260101000000+00'00'"
PDF_TIMESTAMP_PATTERN = re.compile(rb"D:[0-9]{14}(?:[+-][0-9]{2}'[0-9]{2}')?")
NAVY = "20313F"
RED = "B43A46"
TEAL = "167A72"
LIGHT_BLUE = "E8F0F3"
LIGHT_GRAY = "EEF1F3"
MID_GRAY = "B8C0C7"
WHITE = "FFFFFF"


def _normalize_office_zip(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        members = []
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = CORE_TIMESTAMP_PATTERN.sub(
                    lambda match: match.group(1) + CORE_TIMESTAMP + match.group(2), data
                )
            members.append((item.filename, data, item.external_attr))

    file_descriptor, temporary_name = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name, data, external_attr in sorted(members):
                info = zipfile.ZipInfo(name, ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = external_attr
                target.writestr(info, data)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _set_cell_border(cell: object, color: str = MID_GRAY) -> None:
    thin = Side(style="thin", color=color)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)  # type: ignore[attr-defined]


def _format_worksheet(sheet: object, header_row: int, widths: dict[int, float]) -> None:
    sheet.freeze_panes = f"A{header_row + 1}"  # type: ignore[attr-defined]
    sheet.auto_filter.ref = sheet.dimensions  # type: ignore[attr-defined]
    sheet.print_area = sheet.dimensions  # type: ignore[attr-defined]
    sheet.sheet_properties.pageSetUpPr.fitToPage = True  # type: ignore[attr-defined]
    sheet.page_setup.fitToWidth = 2 if max(widths) > 10 else 1  # type: ignore[attr-defined]
    sheet.page_setup.fitToHeight = 1  # type: ignore[attr-defined]
    sheet.page_setup.orientation = "landscape"  # type: ignore[attr-defined]
    sheet.page_setup.paperSize = "8"  # type: ignore[attr-defined]
    sheet.page_margins.left = 0.25  # type: ignore[attr-defined]
    sheet.page_margins.right = 0.25  # type: ignore[attr-defined]
    sheet.page_margins.top = 0.35  # type: ignore[attr-defined]
    sheet.page_margins.bottom = 0.35  # type: ignore[attr-defined]
    sheet.row_dimensions[header_row].height = 34  # type: ignore[attr-defined]
    for cell in sheet[header_row]:  # type: ignore[index]
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        _set_cell_border(cell)
    for row in sheet.iter_rows(min_row=header_row + 1):  # type: ignore[attr-defined]
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            _set_cell_border(cell, "D5DADE")
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width  # type: ignore[attr-defined]


def _new_workbook(title: str, subtitle: str) -> tuple[Workbook, object]:
    workbook = Workbook()
    workbook.properties.creator = "PolicySeal Synthetic Demo"
    workbook.properties.title = title
    workbook.properties.subject = "Synthetic competition fixture"
    workbook.properties.description = "Contains no real company or personal data."
    workbook.properties.created = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)
    sheet = workbook.active
    sheet.title = "Supplier Comparison"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:Q1")
    sheet["A1"] = title
    sheet["A1"].font = Font(size=18, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:Q2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(size=10, italic=True, color="45525C")
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 28
    return workbook, sheet


def _save_workbook(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    _normalize_office_zip(path)


def _create_supplier_matrix_hero(path: Path) -> None:
    workbook, sheet = _new_workbook(
        "FY2026 Supplier Comparison",
        "SYNTHETIC COMPETITION DATA - values and identities are fictional",
    )
    sheet["A3"] = "Decision owner:"
    sheet["B3"] = "Demo Procurement Team"
    sheet["D3"] = "Classification:"
    sheet["E3"] = "Confidential"
    sheet["G3"] = "Review status:"
    sheet["H3"] = "Pending final selection"
    for cell in ("A3", "D3", "G3"):
        sheet[cell].font = Font(bold=True, color=NAVY)
    for cell in ("B3", "E3", "H3"):
        sheet[cell].alignment = Alignment(indent=1)

    headers = [
        "Supplier ID",
        "Supplier Name",
        "Contact Name",
        "Work Email",
        "Mobile",
        "Bank Account",
        "Quoted Price USD",
        "Payment Terms",
        "Internal Cost USD",
        "Target Budget USD",
        "Technical Score",
        "Commercial Score",
        "Rank",
        "Decision Status",
        "Internal Approval Comment",
        "Negotiation Floor USD",
        "Data Owner",
    ]
    sheet.append([])
    sheet.append(headers)
    rows = [
        [
            "SUP-A",
            "Aster Works Ltd.",
            "Maya Chen",
            "maya.chen@aster.example",
            "+1 202 555 0101",
            "9900012345678901",
            482000,
            "Net 45",
            421000,
            500000,
            92,
            88,
            1,
            "Shortlisted",
            "Approved for negotiation; protect the target price.",
            461000,
            "SUP-A",
        ],
        [
            "SUP-B",
            "Blue Peak Systems",
            "Jon Bell",
            "jon.bell@bluepeak.example",
            "+1 202 555 0102",
            "9900012345672244",
            475500,
            "Net 30",
            438000,
            500000,
            86,
            91,
            2,
            "Pending",
            "Legal review pending; do not reveal margin assumptions.",
            468000,
            "SUP-B",
        ],
        [
            "SUP-C",
            "Cedar Ridge Supply",
            "Priya Shah",
            "priya.shah@cedarridge.example",
            "+1 202 555 0103",
            "9900012345677788",
            509000,
            "Net 60",
            449000,
            500000,
            80,
            77,
            3,
            "Not shortlisted",
            "Rejected: delivery risk exceeds the internal threshold.",
            492000,
            "SUP-C",
        ],
    ]
    for row in rows:
        sheet.append(row)
    for row in range(6, 9):
        for column in (7, 9, 10, 16):
            sheet.cell(row, column).number_format = "$#,##0"
        for column in (8, 13, 14, 17):
            sheet.cell(row, column).alignment = Alignment(horizontal="center", vertical="top")
        fill = "F8FAFB" if row % 2 == 0 else WHITE
        for column in range(1, 18):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=fill)
    _format_worksheet(
        sheet,
        header_row=5,
        widths={
            1: 17,
            2: 24,
            3: 18,
            4: 30,
            5: 18,
            6: 22,
            7: 18,
            8: 15,
            9: 18,
            10: 19,
            11: 16,
            12: 18,
            13: 9,
            14: 18,
            15: 42,
            16: 22,
            17: 14,
        },
    )
    glossary = workbook.create_sheet("Field Guide")
    glossary.sheet_view.showGridLines = False
    glossary.append(["Field group", "Synthetic meaning"])
    glossary_rows = [
        ("Supplier and offer", "Fictional supplier-owned business data"),
        ("Internal cost and target", "Demo buyer-only planning data"),
        ("Scores and comments", "Demo procurement decision evidence"),
        ("Contact and account", "Fictional values created only for leakage tests"),
    ]
    for row in glossary_rows:
        glossary.append(row)
    _format_worksheet(glossary, header_row=1, widths={1: 28, 2: 70})
    _save_workbook(workbook, path)


def _create_xlsx_safe(path: Path) -> None:
    workbook, sheet = _new_workbook(
        "Minimal Supplier Offer",
        "SYNTHETIC SAFE EDGE - no personal, bank, or buyer-only fields",
    )
    sheet.append([])
    sheet.append(["Supplier ID", "Supplier Name", "Quoted Price USD", "Data Owner"])
    sheet.append(["SUP-A", "Aster Works Ltd.", 482000, "SUP-A"])
    _format_worksheet(sheet, 5, {1: 16, 2: 28, 3: 22, 4: 16})
    _save_workbook(workbook, path)


def _create_xlsx_unknown(path: Path) -> None:
    workbook, sheet = _new_workbook(
        "Supplier Offer with Unmapped Field",
        "SYNTHETIC APPROVAL EDGE - one field requires controlled mapping",
    )
    sheet.append([])
    sheet.append(
        [
            "Supplier ID",
            "Supplier Name",
            "Quoted Price USD",
            "Strategic Leverage Index",
            "Data Owner",
        ]
    )
    sheet.append(["SUP-A", "Aster Works Ltd.", 482000, "Level-Crimson", "SUP-A"])
    _format_worksheet(sheet, 5, {1: 16, 2: 28, 3: 22, 4: 28, 5: 16})
    _save_workbook(workbook, path)


def _create_xlsx_formula_risk(path: Path) -> None:
    workbook, sheet = _new_workbook(
        "Supplier Formula Risk",
        "SYNTHETIC BLOCK EDGE - formula output cannot be proven safe in P0",
    )
    sheet.append([])
    sheet.append(["Supplier ID", "Quoted Price USD", "Dynamic Internal Margin", "Data Owner"])
    sheet.append(["SUP-A", 482000, "=B6-421000", "SUP-A"])
    _format_worksheet(sheet, 5, {1: 16, 2: 22, 3: 28, 4: 16})
    _save_workbook(workbook, path)


def _write_corrupt_office_file(path: Path, kind: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(
        (
            f"PolicySeal synthetic {kind} processing-failure fixture. "
            "This is intentionally not a ZIP.\n"
        ).encode()
    )


def _shade_cell(cell: object, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _configure_document(document: Document, title: str) -> None:
    properties = document.core_properties
    properties.author = "PolicySeal Synthetic Demo"
    properties.title = title
    properties.subject = "Synthetic competition fixture"
    properties.comments = "Contains no real company or personal data."
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string(NAVY)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor.from_string(NAVY)


def _add_banner(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _shade_cell(cell, LIGHT_BLUE)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def _save_document(document: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _normalize_office_zip(path)


def _create_incident_hero(path: Path) -> None:
    document = Document()
    _configure_document(document, "Synthetic Security Incident Report")
    document.add_heading("Security Incident Report", 0)
    _add_banner(document, "SYNTHETIC COMPETITION DATA - STRICTLY CONFIDENTIAL")
    summary = document.add_table(rows=3, cols=4)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.style = "Table Grid"
    values = [
        ("Incident ID", "INC-DEMO-042", "Status", "Contained"),
        ("Detected", "2026-07-14 08:15 UTC", "Severity", "High"),
        ("Owner", "Demo Security Operations", "Environment", "Synthetic Lab"),
    ]
    for row_index, row_values in enumerate(values):
        for column_index, value in enumerate(row_values):
            cell = summary.cell(row_index, column_index)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if column_index % 2 == 0:
                _shade_cell(cell, LIGHT_GRAY)
                cell.paragraphs[0].runs[0].bold = True

    document.add_heading("1. Executive timeline", level=1)
    document.add_paragraph(
        "At 08:15 UTC the synthetic monitoring service detected repeated archive downloads. "
        "At 08:24 UTC the affected service account was disabled. At 09:10 UTC the demo "
        "environment was isolated, and at 11:40 UTC containment was confirmed."
    )
    document.add_heading("2. Investigation indicators", level=1)
    indicators = document.add_table(rows=1, cols=2)
    indicators.style = "Table Grid"
    indicators.rows[0].cells[0].text = "Indicator type"
    indicators.rows[0].cells[1].text = "Observed synthetic value"
    for cell in indicators.rows[0].cells:
        _shade_cell(cell, NAVY)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].runs[0].bold = True
    for kind, value in [
        ("External IPv4", "203.0.113.77"),
        ("External domain", "dropzone.invalid"),
        ("File SHA-256", "d1f4c3b2a190887766554433221100ffeeddccbbaa99887766554433221100aa"),
        ("Internal host", "fin-app-07.demo.internal / 10.23.8.41"),
    ]:
        cells = indicators.add_row().cells
        cells[0].text = kind
        cells[1].text = value

    document.add_heading("3. Credential and exploit evidence", level=1)
    document.add_paragraph(
        "The lab log contained DEMO_PASSWORD=Crimson-Falcon-47! and token "
        "ps_demo_token_7f4b2e91. These values are fictional leakage-test markers."
    )
    document.add_paragraph(
        "The affected component reported CVE-2025-99999 in Demo Transfer Service 4.2. "
        "Observed behavior was an unauthorized archive request."
    )
    code_paragraph = document.add_paragraph()
    code_run = code_paragraph.add_run(
        "EXPLOIT_DEMO_ONLY: curl -H 'Authorization: Bearer ps_demo_token_7f4b2e91' "
        "https://fin-app-07.demo.internal/export"
    )
    code_run.font.name = "Courier New"
    code_run.font.size = Pt(8.5)
    code_run.font.color.rgb = RGBColor.from_string(RED)

    document.add_heading("4. People and responsibilities", level=1)
    document.add_paragraph(
        "Lead investigator: Elena Park (account epark-admin), Incident Response. "
        "Contact: elena.park@northstar.example, +1 202 555 0144."
    )
    document.add_paragraph(
        "Management note: the delay was attributed to the demo platform team. Performance "
        "follow-up is planned after containment."
    )

    document.add_heading("5. Remediation plan", level=1)
    remediation = document.add_table(rows=1, cols=4)
    remediation.style = "Table Grid"
    remediation.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(("Action", "Owner", "Status / date", "Internal note")):
        remediation.rows[0].cells[index].text = value
        _shade_cell(remediation.rows[0].cells[index], NAVY)
        remediation.rows[0].cells[index].paragraphs[0].runs[0].font.color.rgb = RGBColor(
            255, 255, 255
        )
        remediation.rows[0].cells[index].paragraphs[0].runs[0].bold = True
    for values in [
        ("Rotate service credentials", "Elena Park", "Complete / 2026-07-14", "No external cost"),
        (
            "Patch transfer service",
            "Marcus Lee",
            "In progress / 2026-07-16",
            "Budget cap USD 48,000",
        ),
        (
            "Review archive monitoring",
            "Security Operations",
            "Planned / 2026-07-20",
            "Board note pending",
        ),
    ]:
        cells = remediation.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    document.add_paragraph(
        "End of synthetic incident fixture."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    _save_document(document, path)


def _create_docx_safe(path: Path) -> None:
    document = Document()
    _configure_document(document, "Synthetic Incident Status")
    document.add_heading("Incident Status", 0)
    _add_banner(document, "SYNTHETIC SAFE EDGE")
    document.add_paragraph(
        "Containment completed at 11:40 UTC. The external indicator 203.0.113.77 has been "
        "blocked, and patch validation is in progress."
    )
    _save_document(document, path)


def _create_docx_ambiguous(path: Path) -> None:
    document = Document()
    _configure_document(document, "Synthetic Ambiguous Incident Note")
    document.add_heading("Investigation Note", 0)
    _add_banner(document, "SYNTHETIC APPROVAL EDGE")
    document.add_paragraph(
        "The team agreed that the quiet boundary should remain below the winter threshold. "
        "This phrase has no deterministic category and requires semantic confirmation."
    )
    _save_document(document, path)


def _create_docx_embedded(path: Path) -> None:
    document = Document()
    _configure_document(document, "Synthetic Embedded Object")
    document.add_heading("Incident Package", 0)
    _add_banner(document, "SYNTHETIC BLOCK EDGE - EMBEDDED OBJECT")
    document.add_paragraph("This valid DOCX intentionally contains an unsupported embedded file.")
    _save_document(document, path)
    with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/embeddings/synthetic-demo.bin", b"POLICYSEAL-SYNTHETIC-EMBEDDED")
    _normalize_office_zip(path)


def _pdf_header(page: fitz.Page, title: str, label: str) -> None:
    page.draw_rect(fitz.Rect(0, 0, page.rect.width, 68), color=None, fill=(0.125, 0.192, 0.247))
    page.insert_text((42, 35), title, fontsize=16, fontname="hebo", color=(1, 1, 1))
    page.insert_text((42, 54), label, fontsize=8.5, fontname="helv", color=(0.82, 0.88, 0.91))


def _pdf_footer(page: fitz.Page, page_number: int) -> None:
    page.draw_line(
        fitz.Point(42, page.rect.height - 42),
        fitz.Point(page.rect.width - 42, page.rect.height - 42),
        color=(0.72, 0.75, 0.78),
    )
    page.insert_text(
        (42, page.rect.height - 25),
        "PolicySeal synthetic competition fixture",
        fontsize=7.5,
        fontname="helv",
        color=(0.35, 0.39, 0.43),
    )
    page.insert_text(
        (page.rect.width - 72, page.rect.height - 25),
        f"Page {page_number}",
        fontsize=7.5,
        fontname="helv",
        color=(0.35, 0.39, 0.43),
    )


def _insert_pdf_section(page: fitz.Page, y: float, heading: str, body: str) -> float:
    page.insert_text((42, y), heading, fontsize=11.5, fontname="hebo", color=(0.125, 0.192, 0.247))
    body_rect = fitz.Rect(42, y + 12, page.rect.width - 42, page.rect.height - 58)
    result = page.insert_textbox(
        body_rect,
        body,
        fontsize=9.2,
        fontname="helv",
        lineheight=1.35,
        color=(0.13, 0.15, 0.17),
    )
    used_height = body_rect.height - max(result, 0)
    return y + 18 + used_height + 15


def _save_pdf(document: fitz.Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    for xref in range(1, document.xref_length()):
        if document.xref_get_key(xref, "Type") == ("name", "/EmbeddedFile"):
            document.xref_set_key(xref, "Params/CreationDate", f"({PDF_TIMESTAMP.decode()})")
            document.xref_set_key(xref, "Params/ModDate", f"({PDF_TIMESTAMP.decode()})")
    document.set_metadata(
        {
            "title": "PolicySeal Synthetic Competition Fixture",
            "author": "PolicySeal Synthetic Demo",
            "subject": "Contains no real company or personal data",
            "keywords": "synthetic, competition, policyseal",
            "creator": "PolicySeal",
            "producer": "PolicySeal",
            "creationDate": "D:20260101000000+00'00'",
            "modDate": "D:20260101000000+00'00'",
        }
    )
    document.save(path, garbage=4, deflate=True, no_new_id=True, pretty=False)
    document.close()
    pdf_bytes = path.read_bytes()
    normalized_bytes = PDF_TIMESTAMP_PATTERN.sub(PDF_TIMESTAMP, pdf_bytes)
    if len(normalized_bytes) != len(pdf_bytes):
        raise RuntimeError("PDF timestamp normalization changed file length")
    path.write_bytes(normalized_bytes)


def _create_contract_hero(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    _pdf_header(page, "MASTER SERVICES AGREEMENT", "CONFIDENTIAL - SYNTHETIC SOURCE CONTRACT")
    y = _insert_pdf_section(
        page,
        94,
        "1. Parties and identifiers",
        "Agreement MSA-2024-117 is between Northstar Components Ltd., 18 Meridian Avenue, "
        "Example City, registration 2024-DEMO-771, tax ID TAX-DEMO-4401, represented by "
        "Avery Stone (Buyer), and Redbridge Logistics GmbH, 9 Harbor Lane, Sample Town, "
        "registration 2019-DEMO-118, tax ID TAX-DEMO-9920, represented by Sofia Marin "
        "(Supplier).",
    )
    y = _insert_pdf_section(
        page,
        y,
        "2. Contacts",
        "Buyer contact: Avery Stone, avery.stone@northstar.example, +1 202 555 0160.\n"
        "Supplier contact: Sofia Marin, sofia.marin@redbridge.example, +1 202 555 0161.",
    )
    y = _insert_pdf_section(
        page,
        y,
        "3. Dates and term",
        "Effective date: 1 September 2024. Expiry date: 31 August 2027. Signature date: "
        "20 August 2024. Each delivery must occur within 10 business days after an accepted "
        "purchase order. Notice of delay is due within 2 business days after discovery.",
    )
    _insert_pdf_section(
        page,
        y,
        "4. Historical commercial schedule",
        "Historic unit price: USD 84.25. Volume discount: 7.5%. Custom tax allocation: 3%. "
        "Legacy service credit: 12% of the affected monthly invoice. Standard payment period: "
        "Net 45 days after receipt of an undisputed invoice.",
    )
    _pdf_footer(page, 1)

    page = document.new_page(width=595, height=842)
    _pdf_header(page, "MASTER SERVICES AGREEMENT", "CONFIDENTIAL - SYNTHETIC SOURCE CONTRACT")
    y = _insert_pdf_section(
        page,
        94,
        "5. Standard terms",
        "Each party will comply with applicable written confidentiality obligations. Liability "
        "for ordinary direct damages is limited to fees paid in the preceding twelve months. "
        "Neither party may assign this agreement without prior written consent, except in a "
        "permitted corporate reorganization. Governing law is the fictional law of Example State.",
    )
    y = _insert_pdf_section(
        page,
        y,
        "6. Payment instructions",
        "Payments were historically sent to Demo National Bank, account 8800123456784421, "
        "routing DEMO00077, beneficiary Redbridge Logistics GmbH. These are fictional test values.",
    )
    y = _insert_pdf_section(
        page,
        y,
        "7. Custom service level",
        "NONSTANDARD: Supplier must restore Priority-1 service within 37 minutes and provide a "
        "root-cause report within 19 hours. A service credit of 23% applies for each missed "
        "target.",
    )
    _insert_pdf_section(
        page,
        y,
        "8. Special indemnity",
        "NONSTANDARD: Supplier will indemnify Buyer for all third-party losses arising from any "
        "delay, without the standard aggregate liability cap. This negotiated obligation was "
        "approved only for the original parties.",
    )
    _pdf_footer(page, 2)

    page = document.new_page(width=595, height=842)
    _pdf_header(page, "MASTER SERVICES AGREEMENT", "CONFIDENTIAL - SYNTHETIC SOURCE CONTRACT")
    _insert_pdf_section(
        page,
        94,
        "9. Signatures",
        "For Northstar Components Ltd.\nAvery Stone, Vice President, Procurement\n"
        "Signed 20 August 2024\n\nFor Redbridge Logistics GmbH\n"
        "Sofia Marin, Managing Director\nSigned 20 August 2024\n\n"
        "[SYNTHETIC SEAL: NORTHSTAR]     [SYNTHETIC SEAL: REDBRIDGE]",
    )
    _pdf_footer(page, 3)
    _save_pdf(document, path)


def _create_pdf_text_edge(path: Path, heading: str, body: str, label: str) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    _pdf_header(page, heading, label)
    _insert_pdf_section(page, 105, "Reference terms", body)
    _pdf_footer(page, 1)
    _save_pdf(document, path)


def _create_pdf_attachment(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    _pdf_header(page, "CONTRACT PACKAGE", "SYNTHETIC BLOCK EDGE - EMBEDDED ATTACHMENT")
    _insert_pdf_section(
        page,
        105,
        "Unsupported package",
        "This text page is readable, but the PDF intentionally includes a synthetic attachment.",
    )
    _pdf_footer(page, 1)
    document.embfile_add(
        "synthetic-note.txt",
        b"POLICYSEAL SYNTHETIC EMBEDDED ATTACHMENT",
        filename="synthetic-note.txt",
        desc="Synthetic block fixture",
    )
    _save_pdf(document, path)


def _create_pdf_scanned(path: Path) -> None:
    source = fitz.open()
    source_page = source.new_page(width=595, height=842)
    source_page.draw_rect(source_page.rect, color=None, fill=(1, 1, 1))
    source_page.insert_text(
        (70, 180),
        "SYNTHETIC IMAGE-ONLY CONTRACT PAGE",
        fontsize=18,
        fontname="hebo",
        color=(0.125, 0.192, 0.247),
    )
    source_page.insert_text(
        (70, 220),
        "P0 intentionally does not perform OCR.",
        fontsize=12,
        fontname="helv",
        color=(0.2, 0.22, 0.24),
    )
    pixmap = source_page.get_pixmap(matrix=fitz.Matrix(1, 1), alpha=False)
    image_bytes = pixmap.tobytes("png")
    source.close()

    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=image_bytes)
    _save_pdf(document, path)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def generate_corpus(demo_root: Path) -> dict[str, object]:
    source_root = demo_root / "sources"
    generators: list[tuple[str, str, str, str, str, Callable[[Path], None]]] = [
        (
            "supplier-matrix-hero",
            "sources/supplier-matrix-hero.xlsx",
            "xlsx",
            "hero",
            "safe",
            _create_supplier_matrix_hero,
        ),
        (
            "security-incident-hero",
            "sources/security-incident-hero.docx",
            "docx",
            "hero",
            "safe",
            _create_incident_hero,
        ),
        (
            "reference-contract-hero",
            "sources/reference-contract-hero.pdf",
            "pdf",
            "hero",
            "approval",
            _create_contract_hero,
        ),
        (
            "xlsx-safe-minimal",
            "sources/edges/xlsx-safe-minimal.xlsx",
            "xlsx",
            "edge",
            "safe",
            _create_xlsx_safe,
        ),
        (
            "xlsx-unknown-column",
            "sources/edges/xlsx-unknown-column.xlsx",
            "xlsx",
            "edge",
            "approval",
            _create_xlsx_unknown,
        ),
        (
            "xlsx-formula-risk",
            "sources/edges/xlsx-formula-risk.xlsx",
            "xlsx",
            "edge",
            "block",
            _create_xlsx_formula_risk,
        ),
        (
            "xlsx-corrupt",
            "sources/edges/xlsx-corrupt.xlsx",
            "xlsx",
            "edge",
            "processing_failed",
            lambda path: _write_corrupt_office_file(path, "XLSX"),
        ),
        (
            "docx-safe-minimal",
            "sources/edges/docx-safe-minimal.docx",
            "docx",
            "edge",
            "safe",
            _create_docx_safe,
        ),
        (
            "docx-ambiguous",
            "sources/edges/docx-ambiguous.docx",
            "docx",
            "edge",
            "approval",
            _create_docx_ambiguous,
        ),
        (
            "docx-embedded-object",
            "sources/edges/docx-embedded-object.docx",
            "docx",
            "edge",
            "block",
            _create_docx_embedded,
        ),
        (
            "docx-corrupt",
            "sources/edges/docx-corrupt.docx",
            "docx",
            "edge",
            "processing_failed",
            lambda path: _write_corrupt_office_file(path, "DOCX"),
        ),
        (
            "pdf-safe-standard",
            "sources/edges/pdf-safe-standard.pdf",
            "pdf",
            "edge",
            "safe",
            lambda path: _create_pdf_text_edge(
                path,
                "STANDARD TERMS",
                "Payment is due Net 45 days after an undisputed invoice. Delivery occurs within "
                "10 business days after an accepted order.",
                "SYNTHETIC SAFE EDGE",
            ),
        ),
        (
            "pdf-custom-sla",
            "sources/edges/pdf-custom-sla.pdf",
            "pdf",
            "edge",
            "approval",
            lambda path: _create_pdf_text_edge(
                path,
                "CUSTOM SERVICE LEVEL",
                "NONSTANDARD: restore Priority-1 service within 37 minutes with a 23% service "
                "credit.",
                "SYNTHETIC APPROVAL EDGE",
            ),
        ),
        (
            "pdf-embedded-attachment",
            "sources/edges/pdf-embedded-attachment.pdf",
            "pdf",
            "edge",
            "block",
            _create_pdf_attachment,
        ),
        (
            "pdf-image-only",
            "sources/edges/pdf-image-only.pdf",
            "pdf",
            "edge",
            "block",
            _create_pdf_scanned,
        ),
    ]

    source_root.mkdir(parents=True, exist_ok=True)
    files: list[dict[str, object]] = []
    for source_id, relative_path, kind, role, expected_outcome, generator in generators:
        output_path = demo_root / relative_path
        generator(output_path)
        files.append(
            {
                "id": source_id,
                "path": str(Path("demo_data") / relative_path),
                "kind": kind,
                "role": role,
                "expected_outcome": expected_outcome,
                "size_bytes": output_path.stat().st_size,
                "sha256": _sha256(output_path),
            }
        )

    manifest: dict[str, object] = {
        "schema_version": 1,
        "synthetic_only": True,
        "generator": "scripts/generate_demo_data.py",
        "files": files,
    }
    manifest_path = demo_root / "source_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the frozen PolicySeal synthetic corpus")
    parser.add_argument("--demo-root", type=Path, default=REPO_ROOT / "demo_data")
    args = parser.parse_args()
    manifest = generate_corpus(args.demo_root)
    print(f"Generated {len(manifest['files'])} synthetic files under {args.demo_root}")


if __name__ == "__main__":
    main()
