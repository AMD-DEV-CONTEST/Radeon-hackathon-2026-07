from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from policyseal.application.matrix import default_xlsx_contexts
from policyseal.domain.enums import Channel, RevisionStatus
from policyseal.domain.errors import ExportDenied
from policyseal.domain.models import DisclosureContext


def test_safe_xlsx_persists_stages_and_separate_downloads(
    services,
    hero_xlsx: Path,
) -> None:
    created = services.analysis.create_task(
        hero_xlsx.read_bytes(),
        suffix=".xlsx",
        context=default_xlsx_contexts()[0],
        requester_id="requester-1",
    )

    result = services.analysis.analyze_revision(created.revision.id)

    assert result.revision.status is RevisionStatus.EXPORTABLE
    assert result.revision.output_storage_id is not None
    assert result.events[0].stage == "revision_created"
    assert result.events[-1].status == RevisionStatus.EXPORTABLE.value
    assert {event.stage for event in result.events}.issuperset(
        {
            "preflight",
            "parse",
            "deterministic_scan",
            "semantic_classification",
            "policy_evaluation",
            "render",
            "output_validation",
            "residual_check",
        }
    )

    business = services.exports.business_file(result.revision.id)
    workbook = load_workbook(io.BytesIO(business.data), data_only=False)
    try:
        assert workbook.sheetnames == ["Supplier Comparison", "Field Guide"]
    finally:
        workbook.close()
    evidence = services.exports.evidence_file(result.revision.id)
    with zipfile.ZipFile(io.BytesIO(evidence.data)) as package:
        assert set(package.namelist()) == {
            "decision-report.html",
            "manifest.json",
            "checksums.txt",
        }
        combined = b"\n".join(package.read(name) for name in package.namelist())
    assert b"9900012345678901" not in combined


def test_pdf_approval_blocks_business_download_until_exact_approval(
    services,
    hero_pdf: Path,
) -> None:
    created = services.analysis.create_task(
        hero_pdf.read_bytes(),
        suffix=".pdf",
        context=DisclosureContext.candidate_supplier_corporate_email(),
        requester_id="requester-2",
    )
    result = services.analysis.analyze_revision(created.revision.id)

    assert result.revision.status is RevisionStatus.AWAITING_APPROVAL
    assert len(result.preview_storage_ids) == 3
    with pytest.raises(ExportDenied):
        services.exports.business_file(result.revision.id)
    with pytest.raises(ExportDenied):
        services.exports.evidence_file(result.revision.id)

    approval = services.approvals.request(result.revision.id, actor="requester-2")
    services.approvals.approve(
        result.revision.id,
        actor="demo-reviewer",
        expected_binding_sha256=approval.binding_sha256,
    )
    artifact = services.exports.business_file(
        result.revision.id,
        approval_id=approval.id,
    )
    document = Document(io.BytesIO(artifact.data))
    assert document.paragraphs[0].text == "Reference Draft - Not for Signature"


def test_personal_channel_blocks_and_corrupt_source_fails_closed(
    services,
    hero_xlsx: Path,
) -> None:
    personal_context = default_xlsx_contexts()[2].with_channel(Channel.PERSONAL_EMAIL)
    blocked = services.analysis.create_task(
        hero_xlsx.read_bytes(),
        suffix=".xlsx",
        context=personal_context,
        requester_id="requester-3",
    )
    blocked_result = services.analysis.analyze_revision(blocked.revision.id)

    assert blocked_result.revision.status is RevisionStatus.BLOCKED
    assert blocked_result.revision.output_storage_id is None
    with pytest.raises(ExportDenied):
        services.exports.business_file(blocked_result.revision.id)

    corrupt_path = Path("demo_data/sources/edges/docx-corrupt.docx")
    corrupt = services.analysis.create_task(
        corrupt_path.read_bytes(),
        suffix=".docx",
        context=DisclosureContext.incident_responder_secure_exchange(),
        requester_id="requester-4",
    )
    corrupt_result = services.analysis.analyze_revision(corrupt.revision.id)
    assert corrupt_result.revision.status is RevisionStatus.PROCESSING_FAILED


def test_simulation_is_read_only_and_confirmation_creates_revision_two(
    services,
    hero_xlsx: Path,
) -> None:
    created = services.analysis.create_task(
        hero_xlsx.read_bytes(),
        suffix=".xlsx",
        context=default_xlsx_contexts()[0],
        requester_id="requester-5",
    )
    formal = services.analysis.analyze_revision(created.revision.id)
    simulation = services.simulations.simulate(formal.task.id, default_xlsx_contexts()[1])

    assert simulation.downloadable is False
    assert len(services.repository.list_revisions(formal.task.id)) == 1
    with pytest.raises(ExportDenied):
        services.exports.business_file(simulation.id)

    confirmed = services.simulations.confirm_formal(simulation.id, actor="requester-5")
    revisions = services.repository.list_revisions(formal.task.id)
    assert confirmed.revision.revision_number == 2
    assert confirmed.revision.context == default_xlsx_contexts()[1]
    assert revisions[0].context == default_xlsx_contexts()[0]
    assert revisions[0].status is RevisionStatus.EXPORTABLE
