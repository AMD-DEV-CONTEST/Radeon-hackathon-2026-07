from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from policyseal.application.matrix import default_xlsx_contexts, evaluate_xlsx
from policyseal.documents.preflight import PreflightDisposition
from policyseal.documents.xlsx import XlsxAdapter
from policyseal.domain.enums import ContentCategory
from policyseal.domain.errors import UnsupportedDocumentError
from policyseal.policy.schema import PolicyPack


def _workbook_text(path: Path) -> str:
    workbook = load_workbook(path, data_only=False)
    try:
        return "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()


def test_supplier_matrix_has_stable_regions_and_reopens(
    xlsx_adapter: XlsxAdapter,
    policy_pack: PolicyPack,
    hero_xlsx: Path,
    tmp_path: Path,
) -> None:
    parsed = xlsx_adapter.parse(hero_xlsx)
    repeated = xlsx_adapter.parse(hero_xlsx)
    contexts = default_xlsx_contexts()
    decisions = evaluate_xlsx(parsed, policy_pack, contexts[0])

    assert parsed.regions
    assert len({region.id for region in parsed.regions}) == len(parsed.regions)
    assert tuple(region.id for region in repeated.regions) == tuple(
        region.id for region in parsed.regions
    )
    output = xlsx_adapter.render(parsed, decisions, tmp_path)
    workbook = load_workbook(output.output_path)
    try:
        assert workbook.sheetnames == ["Supplier Comparison", "Field Guide"]
    finally:
        workbook.close()
    assert output.validation.valid


def test_parse_maps_golden_locations_categories_and_owners(
    xlsx_adapter: XlsxAdapter,
    hero_xlsx: Path,
) -> None:
    parsed = xlsx_adapter.parse(hero_xlsx)
    by_location = {region.location: region for region in parsed.regions}

    assert by_location["xlsx:Supplier Comparison:B6"].category is ContentCategory.SUPPLIER_IDENTITY
    assert by_location["xlsx:Supplier Comparison:D6"].category is ContentCategory.CONTACT_PERSONAL
    assert by_location["xlsx:Supplier Comparison:F6"].category is ContentCategory.BANK_PAYMENT
    assert by_location["xlsx:Supplier Comparison:B6"].owner_id == "SUP-A"
    assert by_location["xlsx:Supplier Comparison:B7"].owner_id == "SUP-B"


def test_internal_output_keeps_work_contacts_but_removes_phone_and_bank(
    xlsx_adapter: XlsxAdapter,
    policy_pack: PolicyPack,
    hero_xlsx: Path,
    tmp_path: Path,
) -> None:
    before = hashlib.sha256(hero_xlsx.read_bytes()).hexdigest()
    parsed = xlsx_adapter.parse(hero_xlsx)
    decisions = evaluate_xlsx(parsed, policy_pack, default_xlsx_contexts()[0])

    output = xlsx_adapter.render(parsed, decisions, tmp_path)
    text = _workbook_text(output.output_path)

    assert "Maya Chen" in text
    assert "maya.chen@aster.example" in text
    assert "+1 202 555 0101" not in text
    assert "9900012345678901" not in text
    assert hashlib.sha256(hero_xlsx.read_bytes()).hexdigest() == before


def test_audit_output_keeps_amounts_and_last_four_only(
    xlsx_adapter: XlsxAdapter,
    policy_pack: PolicyPack,
    hero_xlsx: Path,
    tmp_path: Path,
) -> None:
    parsed = xlsx_adapter.parse(hero_xlsx)
    decisions = evaluate_xlsx(parsed, policy_pack, default_xlsx_contexts()[1])

    output = xlsx_adapter.render(parsed, decisions, tmp_path)
    text = _workbook_text(output.output_path)

    assert "482000" in text and "421000" in text
    assert "Account ending 8901" in text
    assert "9900012345678901" not in text
    assert "maya.chen@aster.example" not in text
    assert "protect the target price" not in text
    assert "Approved" in text


def test_bidder_output_contains_only_bidder_a_minimum(
    xlsx_adapter: XlsxAdapter,
    policy_pack: PolicyPack,
    hero_xlsx: Path,
    tmp_path: Path,
) -> None:
    parsed = xlsx_adapter.parse(hero_xlsx)
    decisions = evaluate_xlsx(parsed, policy_pack, default_xlsx_contexts()[2])

    output = xlsx_adapter.render(parsed, decisions, tmp_path)
    text = _workbook_text(output.output_path)

    assert "Aster Works Ltd." in text and "482000" in text and "Shortlisted" in text
    assert "Blue Peak Systems" not in text
    assert "Cedar Ridge Supply" not in text
    assert "maya.chen@aster.example" not in text
    assert "421000" not in text
    assert "protect the target price" not in text


def test_personal_email_block_produces_no_output(
    xlsx_adapter: XlsxAdapter,
    policy_pack: PolicyPack,
    hero_xlsx: Path,
    tmp_path: Path,
) -> None:
    parsed = xlsx_adapter.parse(hero_xlsx)
    decisions = evaluate_xlsx(parsed, policy_pack, default_xlsx_contexts()[3])

    with pytest.raises(UnsupportedDocumentError):
        xlsx_adapter.render(parsed, decisions, tmp_path)

    assert list(tmp_path.iterdir()) == []


def test_unknown_column_requires_confirmation_and_formula_blocks(
    xlsx_adapter: XlsxAdapter,
) -> None:
    unknown = Path("demo_data/sources/edges/xlsx-unknown-column.xlsx")
    formula = Path("demo_data/sources/edges/xlsx-formula-risk.xlsx")

    unknown_result = xlsx_adapter.preflight(unknown)
    formula_result = xlsx_adapter.preflight(formula)

    assert unknown_result.disposition is PreflightDisposition.REQUIRE_CONFIRMATION
    assert formula_result.disposition is PreflightDisposition.BLOCK
    parsed = xlsx_adapter.parse(unknown)
    assert any(region.category is ContentCategory.UNCLASSIFIED for region in parsed.regions)
