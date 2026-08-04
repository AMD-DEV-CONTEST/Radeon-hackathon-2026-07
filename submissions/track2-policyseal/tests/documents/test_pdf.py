from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from docx import Document

from policyseal.documents.pdf import PdfAdapter
from policyseal.documents.preflight import PreflightDisposition
from policyseal.documents.reference_docx import select_safe_nonstandard_placeholders
from policyseal.domain.enums import Action, ContentCategory
from policyseal.domain.models import DisclosureContext
from policyseal.policy.evaluator import evaluate_region
from policyseal.policy.schema import PolicyPack


def _evaluate(
    adapter: PdfAdapter,
    source: Path,
    policy_pack: PolicyPack,
) -> tuple:
    parsed = adapter.parse(source)
    context = DisclosureContext.candidate_supplier_corporate_email()
    return tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)


def _docx_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return "\n".join(values)


def test_scanned_or_embedded_pdf_is_blocked(
    pdf_adapter: PdfAdapter,
    scanned_pdf: Path,
    embedded_pdf: Path,
) -> None:
    scanned = pdf_adapter.preflight(scanned_pdf)
    embedded = pdf_adapter.preflight(embedded_pdf)

    assert scanned.disposition is PreflightDisposition.BLOCK
    assert embedded.disposition is PreflightDisposition.BLOCK
    assert any(finding.code == "pdf_unextractable_page" for finding in scanned.findings)
    assert any(finding.code == "pdf_embedded_file" for finding in embedded.findings)


def test_reference_contract_has_stable_semantic_regions(
    pdf_adapter: PdfAdapter,
    hero_pdf: Path,
) -> None:
    parsed = pdf_adapter.parse(hero_pdf)
    repeated = pdf_adapter.parse(hero_pdf)
    by_location = {region.location: region for region in parsed.regions}

    assert tuple(region.id for region in parsed.regions) == tuple(
        region.id for region in repeated.regions
    )
    assert len(by_location) == len(parsed.regions)
    assert by_location["pdf:page:1:block:1"].category is ContentCategory.CONTRACT_BUYER
    assert by_location["pdf:page:2:block:2"].category is ContentCategory.BANK_PAYMENT
    assert (
        by_location["pdf:page:2:block:3"].category
        is ContentCategory.NONSTANDARD_OBLIGATION
    )
    assert by_location["pdf:page:3:block:1"].category is ContentCategory.SIGNATURE
    assert by_location["pdf:metadata"].category is ContentCategory.SOURCE_METADATA


def test_pdf_preview_renders_local_page_images(
    pdf_adapter: PdfAdapter,
    hero_pdf: Path,
    tmp_path: Path,
) -> None:
    previews = pdf_adapter.preview(hero_pdf, tmp_path)

    assert len(previews) == 3
    assert tuple(item.page for item in previews) == (1, 2, 3)
    assert all(item.path.is_file() and item.path.stat().st_size > 1_000 for item in previews)
    assert all(item.width > 500 and item.height > 700 for item in previews)


def test_reference_contract_requires_approval_and_sanitizes_output(
    pdf_adapter: PdfAdapter,
    policy_pack: PolicyPack,
    hero_pdf: Path,
    tmp_path: Path,
) -> None:
    source_sha256 = hashlib.sha256(hero_pdf.read_bytes()).hexdigest()
    parsed = pdf_adapter.parse(hero_pdf)
    decisions = _evaluate(pdf_adapter, hero_pdf, policy_pack)

    assert any(decision.action is Action.REQUIRE_APPROVAL for decision in decisions)
    result = pdf_adapter.render(parsed, decisions, tmp_path)
    text = _docx_text(result.output_path)
    reopened = Document(result.output_path)

    assert result.output_label == "Reference Draft - Not for Signature"
    assert reopened.paragraphs[0].text == result.output_label
    assert result.output_label in reopened.sections[0].header.paragraphs[0].text
    assert "[Buyer Name]" in text and "[Supplier Name]" in text
    assert "37 minutes" in text
    assert "PolicySeal Traceability" in text
    for forbidden in (
        "Northstar Components Ltd.",
        "Redbridge Logistics GmbH",
        "8800123456784421",
        "Avery Stone",
        "Sofia Marin",
        "[SYNTHETIC SEAL",
        "PolicySeal Synthetic Demo",
    ):
        assert forbidden not in text
    assert hashlib.sha256(hero_pdf.read_bytes()).hexdigest() == source_sha256

    with zipfile.ZipFile(result.output_path) as package:
        names = package.namelist()
        content_types = package.read("[Content_Types].xml")
    assert not any(name.startswith("customXml/") for name in names)
    assert "word/comments.xml" not in names
    assert b"<Types xmlns=" in content_types


def test_safe_nonstandard_placeholder_creates_distinct_output(
    pdf_adapter: PdfAdapter,
    policy_pack: PolicyPack,
    hero_pdf: Path,
    tmp_path: Path,
) -> None:
    parsed = pdf_adapter.parse(hero_pdf)
    exact_decisions = _evaluate(pdf_adapter, hero_pdf, policy_pack)
    safer_decisions = select_safe_nonstandard_placeholders(exact_decisions)

    assert safer_decisions != exact_decisions
    assert not any(
        decision.action is Action.REQUIRE_APPROVAL
        for decision in safer_decisions
        if decision.category is ContentCategory.NONSTANDARD_OBLIGATION
    )
    exact = pdf_adapter.render(parsed, exact_decisions, tmp_path / "exact")
    safer = pdf_adapter.render(parsed, safer_decisions, tmp_path / "safer")
    safer_text = _docx_text(safer.output_path)

    assert exact.output_path.name != safer.output_path.name
    assert "[Nonstandard Obligation - Details Withheld]" in safer_text
    assert "37 minutes" not in safer_text and "23%" not in safer_text
