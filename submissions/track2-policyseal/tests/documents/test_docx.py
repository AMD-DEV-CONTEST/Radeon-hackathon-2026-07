from __future__ import annotations

import hashlib
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from policyseal.documents.docx import DocxAdapter
from policyseal.documents.preflight import PreflightDisposition
from policyseal.domain.enums import Action, ContentCategory
from policyseal.domain.errors import UnsafeOutputError
from policyseal.domain.models import DisclosureContext
from policyseal.policy.evaluator import evaluate_region
from policyseal.policy.schema import PolicyPack


def _decisions(
    adapter: DocxAdapter,
    source: Path,
    policy_pack: PolicyPack,
) -> tuple:
    parsed = adapter.parse(source)
    context = DisclosureContext.incident_responder_secure_exchange()
    return tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)


def _document_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    chunks.extend(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return "\n".join(chunks)


def test_incident_parse_has_stable_golden_regions(
    docx_adapter: DocxAdapter,
    hero_docx: Path,
) -> None:
    parsed = docx_adapter.parse(hero_docx)
    repeated = docx_adapter.parse(hero_docx)
    by_location = {region.location: region for region in parsed.regions}

    assert tuple(region.id for region in parsed.regions) == tuple(
        region.id for region in repeated.regions
    )
    assert len(by_location) == len(parsed.regions)
    assert (
        by_location["docx:section:1:paragraph:1"].category
        is ContentCategory.INCIDENT_FACT_TIMELINE
    )
    assert (
        by_location["docx:section:2:table:1:row:4:cell:2"].category
        is ContentCategory.INTERNAL_INFRASTRUCTURE
    )
    assert (
        by_location["docx:section:3:paragraph:1"].category
        is ContentCategory.CREDENTIAL_SECRET
    )
    assert (
        by_location["docx:section:4:paragraph:1"].category
        is ContentCategory.PERSON_IDENTITY
    )
    assert by_location["docx:section:5:table:1"].category is ContentCategory.REMEDIATION


def test_security_output_reopens_with_aliases_and_without_forbidden_content(
    docx_adapter: DocxAdapter,
    policy_pack: PolicyPack,
    hero_docx: Path,
    tmp_path: Path,
) -> None:
    source_sha256 = hashlib.sha256(hero_docx.read_bytes()).hexdigest()
    parsed = docx_adapter.parse(hero_docx)
    output = docx_adapter.render(
        parsed,
        _decisions(docx_adapter, hero_docx, policy_pack),
        tmp_path,
    )
    text = _document_text(output.output_path)

    assert output.validation.valid
    assert len(Document(output.output_path).tables) == len(Document(hero_docx).tables)
    for required in ("203.0.113.77", "CVE-2025-99999", "Complete / 2026-07-14"):
        assert required in text
    assert "Host-" in text and "Person-" in text and "Account-" in text
    for forbidden in (
        "Crimson-Falcon-47!",
        "ps_demo_token_7f4b2e91",
        "fin-app-07.demo.internal",
        "10.23.8.41",
        "EXPLOIT_DEMO_ONLY",
        "Elena Park",
        "Marcus Lee",
        "epark-admin",
        "elena.park@northstar.example",
        "+1 202 555 0144",
        "Budget cap USD 48,000",
        "Board note pending",
        "delay was attributed",
    ):
        assert forbidden not in text
    assert hashlib.sha256(hero_docx.read_bytes()).hexdigest() == source_sha256

    with zipfile.ZipFile(output.output_path) as package:
        names = package.namelist()
        content_types = package.read("[Content_Types].xml")
        document_relationships = package.read("word/_rels/document.xml.rels")
    assert not any(name.startswith("customXml/") for name in names)
    assert not any(name.startswith("word/embeddings/") for name in names)
    assert b"<Types xmlns=" in content_types
    assert b"<Relationships xmlns=" in document_relationships


def test_aliases_are_stable_across_repeated_outputs(
    docx_adapter: DocxAdapter,
    policy_pack: PolicyPack,
    hero_docx: Path,
    tmp_path: Path,
) -> None:
    parsed = docx_adapter.parse(hero_docx)
    decisions = _decisions(docx_adapter, hero_docx, policy_pack)

    first = _document_text(docx_adapter.render(parsed, decisions, tmp_path / "one").output_path)
    second = _document_text(docx_adapter.render(parsed, decisions, tmp_path / "two").output_path)

    first_aliases = sorted(word for word in first.split() if "-" in word and "-0" in word)
    second_aliases = sorted(word for word in second.split() if "-" in word and "-0" in word)
    assert first_aliases == second_aliases


def test_unresolved_credential_residual_deletes_output(
    docx_adapter: DocxAdapter,
    policy_pack: PolicyPack,
    hero_docx: Path,
    tmp_path: Path,
) -> None:
    parsed = docx_adapter.parse(hero_docx)
    decisions = list(_decisions(docx_adapter, hero_docx, policy_pack))
    credential_index = next(
        index
        for index, decision in enumerate(decisions)
        if decision.category is ContentCategory.CREDENTIAL_SECRET
    )
    decisions[credential_index] = replace(decisions[credential_index], action=Action.KEEP)

    output_dir = tmp_path / "outputs"
    with pytest.raises(UnsafeOutputError):
        docx_adapter.render(parsed, tuple(decisions), output_dir)

    assert list(output_dir.iterdir()) == []


def test_docx_embedded_object_blocks_and_corrupt_file_fails_closed(
    docx_adapter: DocxAdapter,
) -> None:
    embedded = Path("demo_data/sources/edges/docx-embedded-object.docx")
    corrupt = Path("demo_data/sources/edges/docx-corrupt.docx")

    assert docx_adapter.preflight(embedded).disposition is PreflightDisposition.BLOCK
    assert (
        docx_adapter.preflight(corrupt).disposition
        is PreflightDisposition.PROCESSING_FAILED
    )
