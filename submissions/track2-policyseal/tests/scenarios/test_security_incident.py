from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from policyseal.documents.docx import DocxAdapter
from policyseal.domain.enums import Channel
from policyseal.domain.errors import UnsupportedDocumentError
from policyseal.domain.models import DisclosureContext
from policyseal.evals.assets import load_eval_assets
from policyseal.policy.evaluator import evaluate_region
from policyseal.policy.schema import PolicyPack


def _text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return "\n".join(values)


def test_security_incident_matches_independent_golden(
    docx_adapter: DocxAdapter,
    policy_pack: PolicyPack,
    hero_docx: Path,
    tmp_path: Path,
) -> None:
    parsed = docx_adapter.parse(hero_docx)
    context = DisclosureContext.incident_responder_secure_exchange()
    decisions = tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)
    actual = {
        region.location: decision
        for region, decision in zip(parsed.regions, decisions, strict=True)
    }
    golden = next(item for item in load_eval_assets().scenarios if item.id == "scenario-10")

    for expected in golden.decisions:
        decision = actual[expected.location]
        assert decision.action is expected.action
        assert decision.rule_ids == expected.rule_ids
        assert decision.transform == (
            expected.transform.value if expected.transform is not None else None
        )
        assert set(decision.evidence_ids).issuperset(expected.evidence_ids)

    output = docx_adapter.render(parsed, decisions, tmp_path)
    text = _text(output.output_path)
    assert all(literal in text for literal in golden.required_literals)
    assert all(literal not in text for literal in golden.prohibited_literals)


def test_minimal_incident_keeps_timeline_and_external_indicator(
    docx_adapter: DocxAdapter,
    policy_pack: PolicyPack,
    tmp_path: Path,
) -> None:
    source = Path("demo_data/sources/edges/docx-safe-minimal.docx")
    parsed = docx_adapter.parse(source)
    context = DisclosureContext.incident_responder_secure_exchange()
    decisions = tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)
    by_location = {
        region.location: decision
        for region, decision in zip(parsed.regions, decisions, strict=True)
    }

    assert by_location["docx:paragraph:2"].action.value == "keep"
    assert by_location["docx:paragraph:2:indicator"].action.value == "keep"
    output = docx_adapter.render(parsed, decisions, tmp_path)
    text = _text(output.output_path)
    assert "203.0.113.77" in text and "patch validation" in text


def test_incident_personal_email_blocks_without_output(
    docx_adapter: DocxAdapter,
    policy_pack: PolicyPack,
    hero_docx: Path,
    tmp_path: Path,
) -> None:
    parsed = docx_adapter.parse(hero_docx)
    context = DisclosureContext.incident_responder_secure_exchange().with_channel(
        Channel.PERSONAL_EMAIL
    )
    decisions = tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)

    with pytest.raises(UnsupportedDocumentError):
        docx_adapter.render(parsed, decisions, tmp_path)

    assert list(tmp_path.iterdir()) == []
