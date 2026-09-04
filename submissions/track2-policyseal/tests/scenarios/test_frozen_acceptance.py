from __future__ import annotations

import io
import json
import re
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from policyseal.documents.base import DocumentAdapter
from policyseal.domain.enums import Action, ContentCategory, FileKind
from policyseal.domain.errors import ExportDenied, ResidualOutputError
from policyseal.domain.models import DisclosureContext, DocumentRegion
from policyseal.evals.assets import ExpectedStatus, GoldenScenario, load_eval_assets
from policyseal.storage.file_store import FileNamespace

ASSETS = load_eval_assets()
MANIFEST = json.loads(Path("demo_data/source_manifest.json").read_text(encoding="utf-8"))
SOURCES = {item["id"]: Path(item["path"]) for item in MANIFEST["files"]}
pytest_plugins = ("tests.application.conftest",)


class OfflineScenarioClassifier:
    target_location = "docx:section:4:paragraph:2"

    def classify_regions(
        self,
        regions: tuple[DocumentRegion, ...],
    ) -> tuple[DocumentRegion, ...]:
        return tuple(
            replace(
                region,
                category=ContentCategory.UNCLASSIFIED,
                semantic_reason_code="local_model_unavailable",
            )
            if region.location == self.target_location
            else region
            for region in regions
        )


class UnknownOwnerScenarioClassifier:
    def classify_regions(
        self,
        regions: tuple[DocumentRegion, ...],
    ) -> tuple[DocumentRegion, ...]:
        return tuple(
            replace(region, owner_id="__UNKNOWN__")
            if region.location.startswith("xlsx:Supplier Comparison:")
            and region.location[-1:] == "7"
            else region
            for region in regions
        )


class ResidualFailureAdapter:
    file_kind = FileKind.DOCX

    def __init__(self, delegate: DocumentAdapter) -> None:
        self.delegate = delegate

    def preflight(self, path: Path):
        return self.delegate.preflight(path)

    def parse(self, path: Path):
        return self.delegate.parse(path)

    def preview(self, path: Path, output_dir: Path):
        return self.delegate.preview(path, output_dir)

    def reopen_validate(self, path: Path):
        return self.delegate.reopen_validate(path)

    def render(self, parsed, decisions, output_dir):
        del parsed, decisions, output_dir
        raise ResidualOutputError(
            "generated output contains a non-waivable residual risk",
            category=ContentCategory.CREDENTIAL_SECRET,
            location="output:residual-scan:credential",
            rule_id="PS-RESIDUAL-CREDENTIAL",
            evidence_ids=("EV-DOCX-003", "EV-GEN-004"),
        )


def _document_text(data: bytes, suffix: str) -> str:
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            return "\n".join(
                str(cell.value)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
        finally:
            workbook.close()
    document = Document(io.BytesIO(data))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    chunks.extend(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return "\n".join(chunks)


def _contains_literal(text: str, literal: str) -> bool:
    if literal.isdigit():
        return re.search(rf"(?<!\d){re.escape(literal)}(?!\d)", text) is not None
    return literal in text


def _assert_expected_decisions(
    scenario: GoldenScenario,
    regions: tuple[DocumentRegion, ...],
    decisions,
) -> None:
    actual = tuple(zip(regions, decisions, strict=True))
    for expected in scenario.decisions:
        candidates = tuple(
            (region, decision)
            for region, decision in actual
            if region.location == expected.location
        )
        if not candidates and expected.location == "document":
            candidates = actual
        matched = any(
            decision.category is expected.category
            and decision.action is expected.action
            and decision.rule_ids == expected.rule_ids
            and set(decision.evidence_ids).issuperset(expected.evidence_ids)
            and decision.transform
            == (expected.transform.value if expected.transform is not None else None)
            for _, decision in candidates
        )
        details = [
            (
                region.location,
                decision.category.value,
                decision.action.value,
                decision.rule_ids,
                decision.evidence_ids,
            )
            for region, decision in candidates
        ]
        assert matched, (
            f"{scenario.id} did not produce {expected.location} / {expected.rule_ids}; "
            f"candidates={details}"
        )


@pytest.mark.parametrize("scenario", ASSETS.scenarios, ids=lambda item: item.id)
def test_all_frozen_scenarios_reach_their_independent_golden(
    scenario: GoldenScenario,
    services,
    tmp_path: Path,
) -> None:
    source = SOURCES[scenario.source_id]
    if scenario.id == "scenario-15":
        services.analysis.semantic_classifier = UnknownOwnerScenarioClassifier()
    if scenario.id == "scenario-17":
        services.analysis.semantic_classifier = OfflineScenarioClassifier()
    if scenario.id == "scenario-25":
        services.analysis.adapters[FileKind.DOCX] = ResidualFailureAdapter(
            services.analysis.adapters[FileKind.DOCX]
        )

    created = services.analysis.create_task(
        source.read_bytes(),
        suffix=source.suffix,
        context=DisclosureContext(**scenario.context.model_dump()),
        requester_id=f"requester-{scenario.id}",
    )
    result = services.analysis.analyze_revision(created.revision.id)

    assert result.task.source.sha256 == scenario.source_sha256
    assert result.revision.status is scenario.expected_revision_status
    assert all(decision.rule_ids and decision.evidence_ids for decision in result.decisions)
    _assert_expected_decisions(scenario, result.regions, result.decisions)

    if scenario.expected_status is ExpectedStatus.SAFE:
        artifact = services.exports.business_file(result.revision.id)
        output = tmp_path / artifact.filename
        output.write_bytes(artifact.data)
        adapter = services.analysis.adapters[result.task.file_kind]
        assert adapter.reopen_validate(output).valid
        text = _document_text(artifact.data, output.suffix)
        missing = [
            literal
            for literal in scenario.required_literals
            if not _contains_literal(text, literal)
        ]
        leaked = [
            literal for literal in scenario.prohibited_literals if _contains_literal(text, literal)
        ]
        assert not missing, f"{scenario.id} missing required literals: {missing}"
        assert not leaked, f"{scenario.id} retained prohibited literals: {leaked}"
    elif scenario.expected_status is ExpectedStatus.REQUIRE_APPROVAL:
        assert result.revision.output_storage_id is not None
        with pytest.raises(ExportDenied):
            services.exports.business_file(result.revision.id)
        stored = services.file_store.path_for(
            FileNamespace.OUTPUTS,
            result.revision.output_storage_id,
        )
        text = _document_text(stored.read_bytes(), stored.suffix)
        missing = [
            literal
            for literal in scenario.required_literals
            if not _contains_literal(text, literal)
        ]
        leaked = [
            literal for literal in scenario.prohibited_literals if _contains_literal(text, literal)
        ]
        assert not missing, f"{scenario.id} missing required literals: {missing}"
        assert not leaked, f"{scenario.id} retained prohibited literals: {leaked}"
    else:
        assert result.revision.output_storage_id is None
        if scenario.expected_status is ExpectedStatus.BLOCK:
            assert any(decision.action is Action.BLOCK for decision in result.decisions)
        else:
            assert result.decisions == ()
        with pytest.raises(ExportDenied):
            services.exports.business_file(result.revision.id)


def test_frozen_scenario_distribution_matches_the_phase_six_gate() -> None:
    assert ASSETS.scenario_status_counts == {
        "block": 8,
        "processing_failed": 2,
        "require_approval": 8,
        "safe": 12,
    }
