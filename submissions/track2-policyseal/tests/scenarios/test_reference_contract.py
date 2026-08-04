from __future__ import annotations

from pathlib import Path

from docx import Document

from policyseal.documents.pdf import PdfAdapter
from policyseal.domain.models import DisclosureContext
from policyseal.evals.assets import load_eval_assets
from policyseal.policy.evaluator import evaluate_region
from policyseal.policy.schema import PolicyPack


def _text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return "\n".join(values)


def _assert_golden_decisions(parsed, decisions, golden) -> None:
    actual = {
        region.location: decision
        for region, decision in zip(parsed.regions, decisions, strict=True)
    }
    for expected in golden.decisions:
        decision = actual[expected.location]
        assert decision.action is expected.action
        assert decision.rule_ids == expected.rule_ids
        assert decision.transform == (
            expected.transform.value if expected.transform is not None else None
        )
        assert set(decision.evidence_ids).issuperset(expected.evidence_ids)


def test_hero_contract_matches_independent_golden(
    pdf_adapter: PdfAdapter,
    policy_pack: PolicyPack,
    hero_pdf: Path,
    tmp_path: Path,
) -> None:
    parsed = pdf_adapter.parse(hero_pdf)
    context = DisclosureContext.candidate_supplier_corporate_email()
    decisions = tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)
    golden = next(item for item in load_eval_assets().scenarios if item.id == "scenario-18")

    _assert_golden_decisions(parsed, decisions, golden)
    output = pdf_adapter.render(parsed, decisions, tmp_path)
    text = _text(output.output_path)
    assert all(value in text for value in golden.required_literals)
    assert all(value not in text for value in golden.prohibited_literals)


def test_standard_contract_matches_independent_golden(
    pdf_adapter: PdfAdapter,
    policy_pack: PolicyPack,
    tmp_path: Path,
) -> None:
    source = Path("demo_data/sources/edges/pdf-safe-standard.pdf")
    parsed = pdf_adapter.parse(source)
    context = DisclosureContext.candidate_supplier_corporate_email()
    decisions = tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)
    golden = next(item for item in load_eval_assets().scenarios if item.id == "scenario-12")

    _assert_golden_decisions(parsed, decisions, golden)
    output = pdf_adapter.render(parsed, decisions, tmp_path)
    text = _text(output.output_path)
    assert all(value in text for value in golden.required_literals)
    assert all(value not in text for value in golden.prohibited_literals)


def test_compact_custom_sla_matches_independent_golden(
    pdf_adapter: PdfAdapter,
    policy_pack: PolicyPack,
) -> None:
    source = Path("demo_data/sources/edges/pdf-custom-sla.pdf")
    parsed = pdf_adapter.parse(source)
    context = DisclosureContext.candidate_supplier_corporate_email()
    decisions = tuple(evaluate_region(policy_pack, context, region) for region in parsed.regions)
    golden = next(item for item in load_eval_assets().scenarios if item.id == "scenario-19")

    _assert_golden_decisions(parsed, decisions, golden)
