import hashlib
import json
import subprocess
import sys
import time
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook


def _load_manifest(path: Path = Path("demo_data/source_manifest.json")) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def test_frozen_source_manifest_contains_15_valid_files() -> None:
    manifest = _load_manifest()
    files = manifest["files"]
    assert isinstance(files, list)
    assert manifest["synthetic_only"] is True
    assert len(files) == 15
    assert sum(item["role"] == "hero" for item in files) == 3
    assert {item["kind"] for item in files} == {"xlsx", "docx", "pdf"}
    assert all("demo_data/sources/" in item["path"] for item in files)

    for item in files:
        source_path = Path(item["path"])
        data = source_path.read_bytes()
        assert len(data) == item["size_bytes"]
        assert hashlib.sha256(data).hexdigest() == item["sha256"]


def test_hero_files_open_and_contain_expected_synthetic_signals() -> None:
    workbook = load_workbook("demo_data/sources/supplier-matrix-hero.xlsx", data_only=False)
    assert workbook.sheetnames == ["Supplier Comparison", "Field Guide"]
    assert workbook["Supplier Comparison"]["B6"].value == "Aster Works Ltd."

    document = Document("demo_data/sources/security-incident-hero.docx")
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.core_properties.comments == "Contains no real company or personal data."
    assert "DEMO_PASSWORD=Crimson-Falcon-47!" in document_text

    pdf = fitz.open("demo_data/sources/reference-contract-hero.pdf")
    try:
        pdf_text = "\n".join(page.get_text() for page in pdf)
        assert pdf.page_count == 3
        assert pdf.embfile_count() == 0
        assert "Northstar Components Ltd." in pdf_text
        assert "NONSTANDARD" in pdf_text
    finally:
        pdf.close()


def test_generator_is_byte_deterministic(tmp_path: Path) -> None:
    demo_root = tmp_path / "demo_data"
    command = [sys.executable, "scripts/generate_demo_data.py", "--demo-root", str(demo_root)]
    subprocess.run(command, check=True, capture_output=True, text=True)
    first = _load_manifest(demo_root / "source_manifest.json")
    time.sleep(1.1)
    subprocess.run(command, check=True, capture_output=True, text=True)
    second = _load_manifest(demo_root / "source_manifest.json")

    assert first == second
