from __future__ import annotations

import hashlib
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Iterator
from contextlib import suppress
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from policyseal.documents.base import (
    CapabilityFinding,
    OutputValidation,
    ParsedDocument,
    PreflightDisposition,
    PreflightResult,
    PreviewArtifact,
    RenderedDocument,
)
from policyseal.documents.preflight import PreflightService
from policyseal.documents.residual import ResidualScanner, residual_output_error
from policyseal.domain.enums import Action, ContentCategory, FileKind
from policyseal.domain.errors import UnsafeOutputError, UnsupportedDocumentError
from policyseal.domain.models import DocumentRegion, PolicyDecision, SourceFingerprint
from policyseal.security.pseudonyms import (
    PseudonymCategory,
    Pseudonymizer,
)

_SECTION_PATTERN = re.compile(r"^(?P<number>\d+)\.")
_EXTERNAL_IP_PATTERN = re.compile(r"\b(?:203\.0\.113|198\.51\.100|192\.0\.2)\.\d{1,3}\b")
_PERSON_LINE_PATTERN = re.compile(
    r"Lead investigator:\s*(?P<person>[^()]+?)\s*"
    r"\(account\s+(?P<account>[^)]+)\),\s*"
    r"(?P<role>[^.]+)\.\s*(?P<contact>Contact:\s*.+)$",
    re.IGNORECASE,
)


@dataclass(frozen=True, slots=True)
class _DocxTarget:
    location: str
    text: str
    category: ContentCategory
    node: Any | None
    whole_node: bool = True


def _stable_region_id(
    source_sha256: str,
    location: str,
    category: ContentCategory,
) -> str:
    payload = "\0".join((source_sha256, location, category.value)).encode()
    return f"docx-{hashlib.sha256(payload).hexdigest()[:24]}"


def _decision_digest(decisions: tuple[PolicyDecision, ...]) -> str:
    payload = [
        {
            "action": item.action.value,
            "region_id": item.region_id,
            "rule_ids": item.rule_ids,
            "transform": item.transform,
        }
        for item in decisions
    ]
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()[:16]


def _iter_blocks(document: Any) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _paragraph_location(
    paragraph: Paragraph,
    *,
    section: int | None,
    section_paragraph: int,
    body_paragraph: int,
) -> str:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name.startswith("Heading") and section is not None:
        return f"docx:section:{section}:heading"
    if section is not None:
        return f"docx:section:{section}:paragraph:{section_paragraph}"
    return f"docx:paragraph:{body_paragraph}"


def _paragraph_targets(
    paragraph: Paragraph,
    location: str,
    section: int | None,
) -> list[_DocxTarget]:
    text = paragraph.text.strip()
    if not text:
        return []
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name in {"Title", "Subtitle"} or style_name.startswith("Heading"):
        return [_DocxTarget(location, text, ContentCategory.ORGANIZATIONAL_ROLE, paragraph)]

    person_match = _PERSON_LINE_PATTERN.search(text)
    if person_match is not None:
        return [
            _DocxTarget(
                location,
                person_match.group("person").strip(),
                ContentCategory.PERSON_IDENTITY,
                paragraph,
                False,
            ),
            _DocxTarget(
                f"{location}:account",
                person_match.group("account").strip(),
                ContentCategory.INTERNAL_ACCOUNT,
                paragraph,
                False,
            ),
            _DocxTarget(
                f"{location}:role",
                person_match.group("role").strip(),
                ContentCategory.ORGANIZATIONAL_ROLE,
                paragraph,
                False,
            ),
            _DocxTarget(
                f"{location}:contact",
                person_match.group("contact").strip(),
                ContentCategory.CONTACT_DETAIL,
                paragraph,
                False,
            ),
        ]

    normalized = text.casefold()
    if "exploit_demo_only" in normalized or (
        "curl " in normalized and "authorization:" in normalized
    ):
        category = ContentCategory.EXPLOIT_PAYLOAD
    elif any(marker in normalized for marker in ("password=", "token ", "api_key=")):
        category = ContentCategory.CREDENTIAL_SECRET
    elif "management note:" in normalized or "attributed to" in normalized:
        category = ContentCategory.INTERNAL_ATTRIBUTION
    elif "cve-" in normalized:
        category = ContentCategory.VULNERABILITY_FACT
    elif section == 1 or any(
        marker in normalized
        for marker in ("containment", "detected", "isolated", "patch validation")
    ):
        category = ContentCategory.INCIDENT_FACT_TIMELINE
    elif section == 5:
        category = ContentCategory.REMEDIATION
    else:
        category = ContentCategory.UNCLASSIFIED

    targets = [_DocxTarget(location, text, category, paragraph)]
    indicator = _EXTERNAL_IP_PATTERN.search(text)
    if indicator is not None and category is ContentCategory.INCIDENT_FACT_TIMELINE:
        targets.append(
            _DocxTarget(
                f"{location}:indicator",
                indicator.group(0),
                ContentCategory.EXTERNAL_THREAT_INDICATOR,
                paragraph,
                False,
            )
        )
    return targets


def _table_category(
    section: int | None,
    row_index: int,
    column_index: int,
    value: str,
) -> ContentCategory:
    if section == 2:
        if row_index == 4 and column_index == 1:
            return ContentCategory.INTERNAL_INFRASTRUCTURE
        return ContentCategory.EXTERNAL_THREAT_INDICATOR
    if section == 5:
        if row_index > 0 and column_index == 1:
            if "operations" in value.casefold():
                return ContentCategory.ORGANIZATIONAL_ROLE
            return ContentCategory.REMEDIATION_OWNER
        if row_index > 0 and column_index == 3:
            return ContentCategory.REMEDIATION_BUDGET_COMMENT
        return ContentCategory.REMEDIATION
    return ContentCategory.INCIDENT_FACT_TIMELINE


def _table_targets(
    table: Table,
    base_location: str,
    section: int | None,
) -> list[_DocxTarget]:
    if section == 2:
        summary_category = ContentCategory.NETWORK_LOG
    elif section == 5:
        summary_category = ContentCategory.REMEDIATION
    else:
        summary_category = ContentCategory.INCIDENT_FACT_TIMELINE
    targets = [
        _DocxTarget(
            base_location,
            f"{summary_category.value} table structure",
            summary_category,
            None,
        )
    ]
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            value = cell.text.strip()
            if not value:
                continue
            targets.append(
                _DocxTarget(
                    f"{base_location}:row:{row_index}:cell:{column_index + 1}",
                    value,
                    _table_category(section, row_index, column_index, value),
                    cell,
                )
            )
    return targets


def _build_targets(document: Any) -> list[_DocxTarget]:
    targets: list[_DocxTarget] = []
    section: int | None = None
    body_paragraph = 0
    section_paragraph = 0
    body_table = 0
    section_tables: dict[int, int] = {}

    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            body_paragraph += 1
            style_name = block.style.name if block.style is not None else ""
            if style_name.startswith("Heading"):
                match = _SECTION_PATTERN.match(block.text.strip())
                section = int(match.group("number")) if match is not None else (section or 0) + 1
                section_paragraph = 0
            elif section is not None:
                section_paragraph += 1
            location = _paragraph_location(
                block,
                section=section,
                section_paragraph=section_paragraph,
                body_paragraph=body_paragraph,
            )
            targets.extend(_paragraph_targets(block, location, section))
            continue

        if section is None:
            body_table += 1
            table_location = f"docx:table:{body_table}"
        else:
            section_tables[section] = section_tables.get(section, 0) + 1
            table_location = f"docx:section:{section}:table:{section_tables[section]}"
        targets.extend(_table_targets(block, table_location, section))
    return targets


def _replace_target(target: _DocxTarget, replacement: str) -> None:
    if target.node is None:
        raise UnsafeOutputError("a virtual DOCX structure region cannot be transformed")
    current = str(target.node.text)
    if target.whole_node:
        target.node.text = replacement
        return
    if target.text not in current:
        raise UnsafeOutputError("a DOCX fragment no longer matches its stable source region")
    updated = current.replace(target.text, replacement, 1)
    target.node.text = re.sub(r"[ \t]{2,}", " ", updated).strip()


def _copy_sanitized(source: Any, output_path: Path) -> None:
    sanitized = Document()
    source_section = source.sections[0]
    target_section = sanitized.sections[0]
    for attribute in (
        "page_height",
        "page_width",
        "top_margin",
        "right_margin",
        "bottom_margin",
        "left_margin",
    ):
        setattr(target_section, attribute, getattr(source_section, attribute))

    for block in _iter_blocks(source):
        if isinstance(block, Paragraph):
            paragraph = sanitized.add_paragraph()
            style_name = block.style.name if block.style is not None else None
            if style_name is not None:
                with suppress(KeyError):
                    paragraph.style = style_name
            if block.alignment is not None:
                paragraph.alignment = block.alignment
            paragraph.add_run(block.text)
            continue

        column_count = max((len(row.cells) for row in block.rows), default=1)
        copied = sanitized.add_table(rows=len(block.rows), cols=column_count)
        if block.style is not None:
            with suppress(KeyError):
                copied.style = block.style.name
        for row_index, row in enumerate(block.rows):
            for column_index, cell in enumerate(row.cells):
                copied.cell(row_index, column_index).text = cell.text

    properties = sanitized.core_properties
    properties.author = "PolicySeal"
    properties.last_modified_by = "PolicySeal"
    properties.title = "PolicySeal sanitized disclosure"
    properties.subject = ""
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""
    sanitized.save(str(output_path))
    strip_custom_xml(output_path)


def strip_custom_xml(path: Path) -> None:
    temporary_path = path.with_suffix(".sanitizing")
    try:
        with (
            zipfile.ZipFile(path) as source,
            zipfile.ZipFile(
                temporary_path,
                "w",
                compression=zipfile.ZIP_DEFLATED,
            ) as target,
        ):
            for item in source.infolist():
                if item.filename.startswith("customXml/"):
                    continue
                data = source.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    root = ET.fromstring(data)
                    for relationship in list(root):
                        if relationship.attrib.get("Type", "").endswith("/customXml"):
                            root.remove(relationship)
                    ET.register_namespace(
                        "",
                        "http://schemas.openxmlformats.org/package/2006/relationships",
                    )
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif item.filename == "[Content_Types].xml":
                    root = ET.fromstring(data)
                    for declaration in list(root):
                        if declaration.attrib.get("PartName", "").startswith("/customXml/"):
                            root.remove(declaration)
                    ET.register_namespace(
                        "",
                        "http://schemas.openxmlformats.org/package/2006/content-types",
                    )
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(item, data)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


class DocxAdapter:
    file_kind = FileKind.DOCX

    def __init__(
        self,
        preflight: PreflightService,
        residual_scanner: ResidualScanner,
        pseudonymizer: Pseudonymizer,
    ) -> None:
        self._preflight = preflight
        self._residual = residual_scanner
        self._pseudonyms = pseudonymizer

    def preflight(self, path: Path) -> PreflightResult:
        result = self._preflight.inspect(path)
        if result.blocked:
            return result
        try:
            Document(str(path))
        except Exception:
            return PreflightResult(
                file_kind=FileKind.DOCX,
                disposition=PreflightDisposition.PROCESSING_FAILED,
                findings=(
                    CapabilityFinding(
                        code="docx_unreadable_document",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="document",
                        message="The DOCX body cannot be parsed safely.",
                    ),
                ),
            )
        return result

    def parse(self, path: Path) -> ParsedDocument:
        preflight = self.preflight(path)
        if preflight.blocked:
            raise UnsupportedDocumentError("DOCX failed the P0 capability preflight")
        source = SourceFingerprint.from_bytes(path.read_bytes())
        try:
            document = Document(str(path))
            targets = _build_targets(document)
        except Exception as error:
            raise UnsupportedDocumentError("DOCX body cannot be parsed safely") from error
        regions = tuple(
            DocumentRegion(
                id=_stable_region_id(source.sha256, target.location, target.category),
                file_kind=FileKind.DOCX,
                location=target.location,
                order=order,
                text=target.text,
                category=target.category,
            )
            for order, target in enumerate(targets)
        )
        if not regions or len({region.id for region in regions}) != len(regions):
            raise UnsupportedDocumentError("DOCX stable regions are empty or non-unique")
        return ParsedDocument(
            source_path=path,
            source=source,
            file_kind=FileKind.DOCX,
            regions=regions,
            preflight=preflight,
            metadata=(
                ("paragraph_count", str(len(document.paragraphs))),
                ("table_count", str(len(document.tables))),
            ),
        )

    def render(
        self,
        parsed: ParsedDocument,
        decisions: tuple[PolicyDecision, ...],
        output_dir: Path,
    ) -> RenderedDocument:
        if parsed.file_kind is not FileKind.DOCX:
            raise UnsupportedDocumentError("DOCX adapter received a different file kind")
        if any(decision.action is Action.BLOCK for decision in decisions):
            raise UnsupportedDocumentError("blocked disclosure contexts cannot render an output")
        by_region = {decision.region_id: decision for decision in decisions}
        if len(by_region) != len(decisions) or set(by_region) != {
            region.id for region in parsed.regions
        }:
            raise UnsafeOutputError("every stable DOCX region requires exactly one decision")

        document = Document(str(parsed.source_path))
        targets = _build_targets(document)
        target_by_id = {
            _stable_region_id(parsed.source.sha256, target.location, target.category): target
            for target in targets
        }
        if set(target_by_id) != set(by_region):
            raise UnsafeOutputError("DOCX stable regions changed between parse and render")

        scope = parsed.source.sha256
        forbidden_literals: list[str] = []
        for region in parsed.regions:
            decision = by_region[region.id]
            target = target_by_id[region.id]
            if decision.action is Action.KEEP:
                continue
            if decision.action in {Action.REMOVE, Action.REQUIRE_APPROVAL}:
                forbidden_literals.append(target.text)
                _replace_target(target, "")
                continue
            if decision.action is Action.MASK:
                forbidden_literals.append(target.text)
                if decision.transform == "stable_asset_alias":
                    alias_category = PseudonymCategory.HOST
                elif decision.transform == "stable_person_alias":
                    alias_category = (
                        PseudonymCategory.ACCOUNT
                        if target.category is ContentCategory.INTERNAL_ACCOUNT
                        else PseudonymCategory.PERSON
                    )
                else:
                    raise UnsafeOutputError("unsupported DOCX mask transform")
                alias = self._pseudonyms.alias(scope, target.text, alias_category)
                _replace_target(target, alias)
                continue
            raise UnsafeOutputError("unsupported DOCX action")

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / (
            f"policyseal-{parsed.source.sha256[:12]}-{_decision_digest(decisions)}.docx"
        )
        _copy_sanitized(document, output_path)
        validation = self._validate_output(parsed, output_path)
        residual = self._residual.scan(
            output_path,
            forbidden_literals=tuple(dict.fromkeys(forbidden_literals)),
        )
        if not residual.safe:
            output_path.unlink(missing_ok=True)
            raise residual_output_error(residual.findings[0], FileKind.DOCX)
        if not validation.valid:
            output_path.unlink(missing_ok=True)
            raise UnsafeOutputError("generated DOCX failed structure or residual validation")
        return RenderedDocument(output_path, decisions, validation)

    def _validate_output(self, parsed: ParsedDocument, path: Path) -> OutputValidation:
        validation = self.reopen_validate(path)
        if not validation.valid:
            return validation
        expected = dict(parsed.metadata)
        document = Document(str(path))
        findings = list(validation.findings)
        if len(document.paragraphs) != int(expected["paragraph_count"]):
            findings.append("docx_paragraph_count_changed")
        if len(document.tables) != int(expected["table_count"]):
            findings.append("docx_table_count_changed")
        return OutputValidation(not findings, not findings, tuple(findings))

    def reopen_validate(self, path: Path) -> OutputValidation:
        findings: list[str] = []
        try:
            document = Document(str(path))
            has_body = bool(document.paragraphs or document.tables)
            with zipfile.ZipFile(path) as package:
                names = package.namelist()
        except Exception:
            return OutputValidation(False, False, ("docx_reopen_failed",))
        if not has_body:
            findings.append("docx_body_missing")
        if any(name.startswith("customXml/") for name in names):
            findings.append("docx_custom_xml_present")
        if any(name.startswith("word/embeddings/") for name in names):
            findings.append("docx_embedded_content_present")
        if "word/comments.xml" in names:
            findings.append("docx_comments_present")
        return OutputValidation(not findings, not findings, tuple(findings))

    def preview(self, path: Path, output_dir: Path) -> tuple[PreviewArtifact, ...]:
        del path, output_dir
        return ()
