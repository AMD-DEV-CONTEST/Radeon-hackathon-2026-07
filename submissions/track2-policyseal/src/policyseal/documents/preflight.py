from __future__ import annotations

import zipfile
from pathlib import Path

import fitz  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]

from policyseal.documents.base import (
    CapabilityFinding,
    PreflightDisposition,
    PreflightResult,
)
from policyseal.domain.enums import FileKind

_DISPOSITION_RANK = {
    PreflightDisposition.PASS: 0,
    PreflightDisposition.REQUIRE_CONFIRMATION: 1,
    PreflightDisposition.REQUIRE_APPROVAL: 2,
    PreflightDisposition.BLOCK: 3,
    PreflightDisposition.PROCESSING_FAILED: 4,
}


def _result(
    file_kind: FileKind | None,
    findings: list[CapabilityFinding],
) -> PreflightResult:
    disposition = max(
        (finding.disposition for finding in findings),
        key=_DISPOSITION_RANK.__getitem__,
        default=PreflightDisposition.PASS,
    )
    return PreflightResult(file_kind=file_kind, disposition=disposition, findings=tuple(findings))


class PreflightService:
    def inspect(self, path: Path) -> PreflightResult:
        suffix = path.suffix.lower()
        file_kind = {
            ".xlsx": FileKind.XLSX,
            ".docx": FileKind.DOCX,
            ".pdf": FileKind.PDF,
        }.get(suffix)
        if file_kind is None:
            return _result(
                None,
                [
                    CapabilityFinding(
                        code="unsupported_file_kind",
                        disposition=PreflightDisposition.BLOCK,
                        location="file",
                        message="P0 supports only XLSX, DOCX, and text PDF inputs.",
                    )
                ],
            )
        if not path.is_file():
            return _result(
                file_kind,
                [
                    CapabilityFinding(
                        code="source_unavailable",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="file",
                        message="The source file cannot be read.",
                    )
                ],
            )
        if file_kind is FileKind.PDF:
            return self._inspect_pdf(path)
        return self._inspect_office(path, file_kind)

    def _inspect_office(self, path: Path, file_kind: FileKind) -> PreflightResult:
        findings: list[CapabilityFinding] = []
        if not zipfile.is_zipfile(path):
            return _result(
                file_kind,
                [
                    CapabilityFinding(
                        code=f"{file_kind.value}_invalid_package",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="package",
                        message="The Office package is corrupt or unreadable.",
                    )
                ],
            )
        try:
            with zipfile.ZipFile(path) as package:
                names = tuple(package.namelist())
        except (OSError, zipfile.BadZipFile):
            return _result(
                file_kind,
                [
                    CapabilityFinding(
                        code=f"{file_kind.value}_invalid_package",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="package",
                        message="The Office package cannot be inspected.",
                    )
                ],
            )

        embedded_prefixes = (
            "word/embeddings/",
            "xl/embeddings/",
            "xl/externalLinks/",
        )
        if any(name.startswith(embedded_prefixes) for name in names):
            location = (
                "document:embedded-object"
                if file_kind is FileKind.DOCX
                else "workbook:embedded-content"
            )
            findings.append(
                CapabilityFinding(
                    code=f"{file_kind.value}_embedded_content",
                    disposition=PreflightDisposition.BLOCK,
                    location=location,
                    message="Embedded or externally linked content is unsupported in P0.",
                )
            )
        if file_kind is FileKind.XLSX:
            try:
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    formula_location = next(
                        (
                            f"xlsx:{sheet.title}:{cell.coordinate}"
                            for sheet in workbook.worksheets
                            for row in sheet.iter_rows()
                            for cell in row
                            if isinstance(cell.value, str) and cell.value.startswith("=")
                        ),
                        None,
                    )
                finally:
                    workbook.close()
            except Exception:
                findings.append(
                    CapabilityFinding(
                        code="xlsx_unreadable_workbook",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="workbook",
                        message="The workbook cannot be parsed safely.",
                    )
                )
            else:
                if formula_location is not None:
                    findings.append(
                        CapabilityFinding(
                            code="xlsx_formula_unsupported",
                            disposition=PreflightDisposition.BLOCK,
                            location=formula_location,
                            message="Formula evaluation is outside the P0 safety boundary.",
                        )
                    )
        return _result(file_kind, findings)

    def _inspect_pdf(self, path: Path) -> PreflightResult:
        findings: list[CapabilityFinding] = []
        try:
            document = fitz.open(path)
        except Exception:
            return _result(
                FileKind.PDF,
                [
                    CapabilityFinding(
                        code="pdf_unreadable",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="pdf",
                        message="The PDF cannot be opened.",
                    )
                ],
            )
        try:
            if document.needs_pass:
                findings.append(
                    CapabilityFinding(
                        code="pdf_encrypted",
                        disposition=PreflightDisposition.BLOCK,
                        location="pdf",
                        message="Encrypted PDFs are unsupported in P0.",
                    )
                )
            if document.embfile_count() > 0:
                findings.append(
                    CapabilityFinding(
                        code="pdf_embedded_file",
                        disposition=PreflightDisposition.BLOCK,
                        location="document:embedded-attachment",
                        message="A PDF package with embedded files cannot be safely transformed.",
                    )
                )
            if document.page_count == 0:
                findings.append(
                    CapabilityFinding(
                        code="pdf_empty",
                        disposition=PreflightDisposition.BLOCK,
                        location="pdf",
                        message="A PDF with no pages cannot be processed.",
                    )
                )
            for page_index, page in enumerate(document, start=1):
                if not page.get_text("text").strip():
                    findings.append(
                        CapabilityFinding(
                            code="pdf_unextractable_page",
                            disposition=PreflightDisposition.BLOCK,
                            location=f"pdf:page:{page_index}",
                            message="The page has no reliable text layer; OCR is outside P0.",
                        )
                    )
                if page.first_annot is not None:
                    findings.append(
                        CapabilityFinding(
                            code="pdf_annotations_present",
                            disposition=PreflightDisposition.REQUIRE_APPROVAL,
                            location=f"pdf:page:{page_index}",
                            message="PDF annotations are excluded from generated output.",
                        )
                    )
        finally:
            document.close()
        return _result(FileKind.PDF, findings)


__all__ = ["PreflightDisposition", "PreflightResult", "PreflightService"]
