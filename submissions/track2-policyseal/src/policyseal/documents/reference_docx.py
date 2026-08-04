from __future__ import annotations

import re
from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from policyseal.documents.base import ParsedDocument
from policyseal.documents.docx import strip_custom_xml
from policyseal.domain.enums import Action, ContentCategory, DecisionOrigin
from policyseal.domain.errors import UnsafeOutputError
from policyseal.domain.models import PolicyDecision

REFERENCE_DRAFT_LABEL = "Reference Draft - Not for Signature"
_SAFE_NONSTANDARD_TRANSFORM = "placeholder_nonstandard_obligation"


def select_safe_nonstandard_placeholders(
    decisions: tuple[PolicyDecision, ...],
) -> tuple[PolicyDecision, ...]:
    return tuple(
        replace(
            decision,
            action=Action.GENERALIZE,
            transform=_SAFE_NONSTANDARD_TRANSFORM,
            origin=DecisionOrigin.HUMAN_CONFIRMED,
            reason_code="safer_nonstandard_placeholder_selected",
            reason="The requester selected a stricter placeholder instead of exact disclosure.",
        )
        if decision.category is ContentCategory.NONSTANDARD_OBLIGATION
        and decision.action is Action.REQUIRE_APPROVAL
        else decision
        for decision in decisions
    )


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _placeholder(region_text: str, transform: str | None) -> str:
    if transform == "placeholder_buyer":
        return "[Buyer Name]"
    if transform == "placeholder_supplier":
        return "[Supplier Name]"
    if transform == "placeholder_corporate_id":
        role = "Buyer" if "(Buyer)" in region_text else "Supplier"
        return f"[{role} Registered Address and Corporate Identifiers]"
    if transform == "placeholder_contact":
        role = "Buyer" if region_text.casefold().startswith("buyer") else "Supplier"
        return f"[{role} Contact]"
    if transform == "placeholder_contract_date":
        if re.fullmatch(r"[A-Z]{2,}-[A-Z0-9-]+", region_text):
            return "[Contract Identifier]"
        return "[Contract Date or Absolute Date]"
    if transform == "placeholder_historical_pricing":
        return "[Historical Commercial Values]"
    if transform == _SAFE_NONSTANDARD_TRANSFORM:
        return "[Nonstandard Obligation - Details Withheld]"
    raise UnsafeOutputError("unsupported reference-DOCX placeholder transform")


def _rendered_text(region_text: str, decision: PolicyDecision) -> str | None:
    if decision.action is Action.REMOVE:
        return None
    if decision.action is Action.KEEP:
        return _normalize_text(region_text)
    if decision.action is Action.REQUIRE_APPROVAL:
        return (
            _normalize_text(region_text)
            if decision.category is ContentCategory.NONSTANDARD_OBLIGATION
            else None
        )
    if decision.action is Action.GENERALIZE:
        return _placeholder(region_text, decision.transform)
    raise UnsafeOutputError("unsupported action for a PDF reference draft")


def _page_number(location: str) -> int | None:
    match = re.match(r"pdf:page:(\d+):", location)
    return int(match.group(1)) if match is not None else None


def render_reference_docx(
    parsed: ParsedDocument,
    decisions: tuple[PolicyDecision, ...],
    output_path: Path,
) -> None:
    by_region = {decision.region_id: decision for decision in decisions}
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.right_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(REFERENCE_DRAFT_LABEL)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Candidate Supplier Contract Terms Preview")
    subtitle_run.bold = True

    if any(decision.action is Action.REQUIRE_APPROVAL for decision in decisions):
        notice = document.add_paragraph()
        notice_run = notice.add_run(
            "Approval required: exact nonstandard obligations may be included for review; "
            "other unapproved content is withheld."
        )
        notice_run.bold = True

    current_page: int | None = None
    rendered_count = 0
    for region in parsed.regions:
        decision = by_region[region.id]
        if region.category is ContentCategory.SOURCE_METADATA:
            continue
        text = _rendered_text(region.text, decision)
        if text is None:
            continue
        page = _page_number(region.location)
        if page is not None and page != current_page:
            document.add_heading(f"Source Page {page}", level=1)
            current_page = page
        paragraph = document.add_paragraph(text)
        if decision.action is Action.REQUIRE_APPROVAL:
            paragraph.style = "Quote"
            paragraph.add_run(" [Approval Required]").bold = True
        rendered_count += 1
    if rendered_count == 0:
        raise UnsafeOutputError("reference draft has no renderable contract content")

    document.add_page_break()  # type: ignore[no-untyped-call]
    document.add_heading("PolicySeal Traceability", level=1)
    trace = document.add_table(rows=1, cols=3)
    trace.style = "Table Grid"
    headers = trace.rows[0].cells
    headers[0].text = "Source region"
    headers[1].text = "Action"
    headers[2].text = "Rule evidence"
    for region in parsed.regions:
        decision = by_region[region.id]
        cells = trace.add_row().cells
        cells[0].text = region.location
        cells[1].text = decision.action.value
        cells[2].text = ", ".join((*decision.rule_ids, *decision.evidence_ids))

    header = section.header.paragraphs[0]
    header.text = REFERENCE_DRAFT_LABEL
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.text = f"PolicySeal source checksum {parsed.source.sha256[:16]}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)

    properties = document.core_properties
    properties.author = "PolicySeal"
    properties.last_modified_by = "PolicySeal"
    properties.title = REFERENCE_DRAFT_LABEL
    properties.subject = "Candidate supplier contract terms preview"
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""
    document.save(str(output_path))
    strip_custom_xml(output_path)
