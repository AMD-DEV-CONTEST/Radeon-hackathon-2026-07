from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import fitz  # type: ignore[import-untyped]
from docx import Document

from policyseal.documents.base import (
    OutputValidation,
    ParsedDocument,
    PreflightResult,
    PreviewArtifact,
    RenderedDocument,
)
from policyseal.documents.preflight import PreflightService
from policyseal.documents.reference_docx import (
    REFERENCE_DRAFT_LABEL,
    render_reference_docx,
)
from policyseal.documents.residual import ResidualScanner, residual_output_error
from policyseal.domain.enums import Action, ContentCategory, FileKind
from policyseal.domain.errors import UnsafeOutputError, UnsupportedDocumentError
from policyseal.domain.models import DocumentRegion, PolicyDecision, SourceFingerprint

_SECTION_HEADING = re.compile(r"^\d+\.\s+[^.\n]{2,80}$")
_SENTENCE_BOUNDARY = re.compile(r"(?<=[.!?])\s+(?=[A-Z])")


@dataclass(frozen=True, slots=True)
class _SemanticFragment:
    text: str
    category: ContentCategory
    bbox: tuple[float, float, float, float]


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _stable_region_id(
    source_sha256: str,
    location: str,
    category: ContentCategory,
) -> str:
    payload = "\0".join((source_sha256, location, category.value)).encode()
    return f"pdf-{hashlib.sha256(payload).hexdigest()[:24]}"


def _decision_digest(decisions: tuple[PolicyDecision, ...]) -> str:
    payload = [
        {
            "action": item.action.value,
            "region_id": item.region_id,
            "rule_ids": item.rule_ids,
            "transform": item.transform,
        }
        for item in decisions
    ]
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()[:16]


def _is_structural(text: str, y0: float, page_height: float) -> bool:
    normalized = _normalize(text)
    if not normalized or y0 > page_height - 60:
        return True
    if normalized.casefold() == "reference terms":
        return True
    if _SECTION_HEADING.fullmatch(normalized):
        return True
    letters = "".join(character for character in normalized if character.isalpha())
    return bool(letters) and letters.upper() == letters and len(normalized) < 160


def _sentences(text: str) -> tuple[str, ...]:
    return tuple(part.strip() for part in _SENTENCE_BOUNDARY.split(text) if part.strip())


def _party_fragments(
    text: str,
    bbox: tuple[float, float, float, float],
) -> list[_SemanticFragment]:
    match = re.match(r"Agreement\s+(\S+)\s+is between\s+(.+)\.$", text)
    if match is None or ", and " not in match.group(2):
        return [_SemanticFragment(text, ContentCategory.UNCLASSIFIED, bbox)]
    agreement_id = match.group(1)
    buyer_segment, supplier_segment = match.group(2).split(", and ", 1)
    if "," not in buyer_segment or "," not in supplier_segment:
        return [_SemanticFragment(text, ContentCategory.UNCLASSIFIED, bbox)]
    buyer_name, buyer_details = buyer_segment.split(",", 1)
    supplier_name, supplier_details = supplier_segment.split(",", 1)
    return [
        _SemanticFragment(buyer_name.strip(), ContentCategory.CONTRACT_BUYER, bbox),
        _SemanticFragment(supplier_name.strip(), ContentCategory.CONTRACT_SUPPLIER, bbox),
        _SemanticFragment(buyer_details.strip(), ContentCategory.CORPORATE_IDENTIFIER, bbox),
        _SemanticFragment(supplier_details.strip(), ContentCategory.CORPORATE_IDENTIFIER, bbox),
        _SemanticFragment(agreement_id, ContentCategory.CONTRACT_DATE_IDENTIFIER, bbox),
    ]


def _contact_fragments(
    text: str,
    bbox: tuple[float, float, float, float],
) -> list[_SemanticFragment]:
    parts = re.split(r"(?=Supplier contact:)", text)
    return [
        _SemanticFragment(part.strip(), ContentCategory.CONTACT_DETAIL, bbox)
        for part in parts
        if part.strip()
    ]


def _date_fragments(
    text: str,
    bbox: tuple[float, float, float, float],
) -> list[_SemanticFragment]:
    fragments: list[_SemanticFragment] = []
    for sentence in _sentences(text):
        category = (
            ContentCategory.CONTRACT_DATE_IDENTIFIER
            if sentence.startswith(("Effective date:", "Expiry date:", "Signature date:"))
            else ContentCategory.RELATIVE_PERFORMANCE_TERM
        )
        fragments.append(_SemanticFragment(sentence, category, bbox))
    return fragments


def _historical_fragments(
    text: str,
    bbox: tuple[float, float, float, float],
) -> list[_SemanticFragment]:
    marker = "Standard payment period:"
    if marker not in text:
        return [_SemanticFragment(text, ContentCategory.HISTORICAL_PRICING, bbox)]
    historical, payment = text.split(marker, 1)
    return [
        _SemanticFragment(historical.strip(), ContentCategory.HISTORICAL_PRICING, bbox),
        _SemanticFragment(
            f"{marker} {payment.strip()}",
            ContentCategory.STANDARD_LEGAL_TERM,
            bbox,
        ),
    ]


def _generic_fragments(
    text: str,
    bbox: tuple[float, float, float, float],
) -> list[_SemanticFragment]:
    fragments: list[_SemanticFragment] = []
    for sentence in _sentences(text):
        normalized = sentence.casefold()
        if "nonstandard:" in normalized or "indemnify" in normalized:
            category = ContentCategory.NONSTANDARD_OBLIGATION
        elif "account " in normalized and ("bank" in normalized or "routing" in normalized):
            category = ContentCategory.BANK_PAYMENT
        elif any(
            marker in normalized
            for marker in (
                "payment is due",
                "standard payment period",
                "confidentiality obligations",
                "liability for ordinary direct damages",
                "governing law",
                "neither party may assign",
            )
        ):
            category = ContentCategory.STANDARD_LEGAL_TERM
        elif re.search(r"\bwithin\s+\d+\s+(?:business\s+)?(?:days|hours|minutes)\b", normalized):
            category = ContentCategory.RELATIVE_PERFORMANCE_TERM
        else:
            category = ContentCategory.UNCLASSIFIED
        fragments.append(_SemanticFragment(sentence, category, bbox))
    return fragments


def _semantic_fragments(
    text: str,
    bbox: tuple[float, float, float, float],
) -> list[_SemanticFragment]:
    if text.startswith("Agreement ") and " is between " in text:
        return _party_fragments(text, bbox)
    if text.startswith("Buyer contact:"):
        return _contact_fragments(text, bbox)
    if text.startswith("Effective date:"):
        return _date_fragments(text, bbox)
    if text.startswith("Historic unit price:"):
        return _historical_fragments(text, bbox)
    if text.startswith("Each party will comply"):
        return [_SemanticFragment(text, ContentCategory.STANDARD_LEGAL_TERM, bbox)]
    if text.startswith("Payments were historically sent"):
        return [_SemanticFragment(text, ContentCategory.BANK_PAYMENT, bbox)]
    if text.startswith("NONSTANDARD:"):
        return [_SemanticFragment(text, ContentCategory.NONSTANDARD_OBLIGATION, bbox)]
    return _generic_fragments(text, bbox)


def _page_content_blocks(page: fitz.Page) -> list[tuple[str, tuple[float, float, float, float]]]:
    blocks: list[tuple[str, tuple[float, float, float, float]]] = []
    for block in sorted(page.get_text("blocks"), key=lambda item: (item[1], item[0])):
        if len(block) < 7 or block[6] != 0:
            continue
        text = _normalize(str(block[4]))
        if _is_structural(text, float(block[1]), float(page.rect.height)):
            continue
        blocks.append(
            (
                text,
                (float(block[0]), float(block[1]), float(block[2]), float(block[3])),
            )
        )
    return blocks


def _extract_regions(
    document: fitz.Document,
    source_sha256: str,
) -> tuple[tuple[DocumentRegion, ...], tuple[tuple[str, str], ...]]:
    regions: list[DocumentRegion] = []
    metadata: list[tuple[str, str]] = [("page_count", str(document.page_count))]
    for page_number, page in enumerate(document, start=1):
        blocks = _page_content_blocks(page)
        fragments: list[_SemanticFragment] = []
        page_text = _normalize(page.get_text("text"))
        if "9. Signatures" in page_text and blocks:
            combined = " ".join(text for text, _ in blocks)
            x0 = min(bbox[0] for _, bbox in blocks)
            y0 = min(bbox[1] for _, bbox in blocks)
            x1 = max(bbox[2] for _, bbox in blocks)
            y1 = max(bbox[3] for _, bbox in blocks)
            fragments.append(
                _SemanticFragment(combined, ContentCategory.SIGNATURE, (x0, y0, x1, y1))
            )
        else:
            for text, bbox in blocks:
                fragments.extend(_semantic_fragments(text, bbox))
        for block_number, fragment in enumerate(fragments, start=1):
            location = f"pdf:page:{page_number}:block:{block_number}"
            regions.append(
                DocumentRegion(
                    id=_stable_region_id(source_sha256, location, fragment.category),
                    file_kind=FileKind.PDF,
                    location=location,
                    order=len(regions),
                    text=fragment.text,
                    category=fragment.category,
                )
            )
            metadata.append(
                (
                    f"bbox:{location}",
                    ",".join(f"{coordinate:.2f}" for coordinate in fragment.bbox),
                )
            )
        metadata.append((f"images:page:{page_number}", str(len(page.get_images(full=True)))))

    metadata_location = "pdf:metadata"
    regions.append(
        DocumentRegion(
            id=_stable_region_id(
                source_sha256,
                metadata_location,
                ContentCategory.SOURCE_METADATA,
            ),
            file_kind=FileKind.PDF,
            location=metadata_location,
            order=len(regions),
            text="Source PDF metadata and annotations",
            category=ContentCategory.SOURCE_METADATA,
        )
    )
    return tuple(regions), tuple(metadata)


class PdfAdapter:
    file_kind = FileKind.PDF

    def __init__(self, preflight: PreflightService, residual_scanner: ResidualScanner) -> None:
        self._preflight = preflight
        self._residual = residual_scanner

    def preflight(self, path: Path) -> PreflightResult:
        return self._preflight.inspect(path)

    def parse(self, path: Path) -> ParsedDocument:
        preflight = self.preflight(path)
        if preflight.blocked:
            raise UnsupportedDocumentError("PDF failed the P0 capability preflight")
        source = SourceFingerprint.from_bytes(path.read_bytes())
        try:
            document = fitz.open(path)
            try:
                regions, metadata = _extract_regions(document, source.sha256)
            finally:
                document.close()
        except Exception as error:
            raise UnsupportedDocumentError("PDF text cannot be mapped safely") from error
        if len(regions) <= 1 or len({region.id for region in regions}) != len(regions):
            raise UnsupportedDocumentError("PDF semantic regions are empty or non-unique")
        return ParsedDocument(path, source, FileKind.PDF, regions, preflight, metadata)

    def render(
        self,
        parsed: ParsedDocument,
        decisions: tuple[PolicyDecision, ...],
        output_dir: Path,
    ) -> RenderedDocument:
        if parsed.file_kind is not FileKind.PDF:
            raise UnsupportedDocumentError("PDF adapter received a different file kind")
        if any(decision.action is Action.BLOCK for decision in decisions):
            raise UnsupportedDocumentError("blocked disclosure contexts cannot render an output")
        by_region = {decision.region_id: decision for decision in decisions}
        if len(by_region) != len(decisions) or set(by_region) != {
            region.id for region in parsed.regions
        }:
            raise UnsafeOutputError("every stable PDF region requires exactly one decision")

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / (
            f"policyseal-{parsed.source.sha256[:12]}-{_decision_digest(decisions)}.docx"
        )
        render_reference_docx(parsed, decisions, output_path)
        validation = self.reopen_validate(output_path)
        forbidden_literals = tuple(
            region.text
            for region in parsed.regions
            if by_region[region.id].action in {Action.REMOVE, Action.GENERALIZE}
            or (
                by_region[region.id].action is Action.REQUIRE_APPROVAL
                and by_region[region.id].category is not ContentCategory.NONSTANDARD_OBLIGATION
            )
        )
        residual = self._residual.scan(output_path, forbidden_literals=forbidden_literals)
        if not residual.safe:
            output_path.unlink(missing_ok=True)
            raise residual_output_error(residual.findings[0], FileKind.PDF)
        if not validation.valid:
            output_path.unlink(missing_ok=True)
            raise UnsafeOutputError("reference draft failed structure or residual validation")
        return RenderedDocument(
            output_path,
            decisions,
            validation,
            output_label=REFERENCE_DRAFT_LABEL,
        )

    def reopen_validate(self, path: Path) -> OutputValidation:
        findings: list[str] = []
        try:
            document = Document(str(path))
            first_text = document.paragraphs[0].text if document.paragraphs else ""
            header_text = "\n".join(
                paragraph.text
                for section in document.sections
                for paragraph in section.header.paragraphs
            )
            with zipfile.ZipFile(path) as package:
                names = package.namelist()
        except Exception:
            return OutputValidation(False, False, ("reference_docx_reopen_failed",))
        if first_text != REFERENCE_DRAFT_LABEL:
            findings.append("reference_label_missing")
        if REFERENCE_DRAFT_LABEL not in header_text:
            findings.append("reference_header_missing")
        if any(name.startswith("customXml/") for name in names):
            findings.append("reference_custom_xml_present")
        if "word/comments.xml" in names:
            findings.append("reference_comments_present")
        return OutputValidation(not findings, not findings, tuple(findings))

    def preview(self, path: Path, output_dir: Path) -> tuple[PreviewArtifact, ...]:
        preflight = self.preflight(path)
        if preflight.blocked:
            raise UnsupportedDocumentError("blocked PDFs cannot produce a preview")
        source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
        output_dir.mkdir(parents=True, exist_ok=True)
        document = fitz.open(path)
        previews: list[PreviewArtifact] = []
        try:
            for page_number, page in enumerate(document, start=1):
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                preview_path = output_dir / f"{source_sha256[:12]}-page-{page_number}.png"
                pixmap.save(preview_path)
                previews.append(
                    PreviewArtifact(
                        page=page_number,
                        path=preview_path,
                        width=pixmap.width,
                        height=pixmap.height,
                    )
                )
        finally:
            document.close()
        return tuple(previews)
