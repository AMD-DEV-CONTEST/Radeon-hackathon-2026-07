from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

import fitz  # type: ignore[import-untyped]
from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]

from policyseal.domain.enums import ContentCategory, FileKind
from policyseal.domain.errors import ResidualOutputError
from policyseal.security.fingerprints import KeyedFingerprinter
from policyseal.security.key_store import MasterKeyStore


@dataclass(frozen=True, slots=True)
class ResidualFinding:
    pattern_id: str
    category: ContentCategory
    location: str
    value_hmac: str
    waivable: bool


@dataclass(frozen=True, slots=True)
class ResidualScanResult:
    safe: bool
    waivable: bool
    output_sha256: str
    findings: tuple[ResidualFinding, ...]


@dataclass(frozen=True, slots=True)
class _HighRiskPattern:
    id: str
    category: ContentCategory
    expression: re.Pattern[str]


_PATTERNS = (
    _HighRiskPattern(
        "credential_assignment",
        ContentCategory.CREDENTIAL_SECRET,
        re.compile(r"(?i)\b(?:password|passwd|api[_-]?key|secret)\s*[:=]\s*[^\s,;]+"),
    ),
    _HighRiskPattern(
        "policyseal_demo_token",
        ContentCategory.CREDENTIAL_SECRET,
        re.compile(r"\bps_demo_token_[A-Za-z0-9_\-]+\b"),
    ),
    _HighRiskPattern(
        "aws_access_key",
        ContentCategory.CREDENTIAL_SECRET,
        re.compile(r"\bAKIA[0-9A-Z]{16}\b"),
    ),
    _HighRiskPattern(
        "bearer_token",
        ContentCategory.CREDENTIAL_SECRET,
        re.compile(r"(?i)\bBearer\s+[A-Za-z0-9._~+/=\-]{8,}"),
    ),
    _HighRiskPattern(
        "private_key",
        ContentCategory.CREDENTIAL_SECRET,
        re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----"),
    ),
    _HighRiskPattern(
        "full_payment_account",
        ContentCategory.BANK_PAYMENT,
        re.compile(r"(?<!\d)\d{12,19}(?!\d)"),
    ),
    _HighRiskPattern(
        "exploit_command",
        ContentCategory.EXPLOIT_PAYLOAD,
        re.compile(r"(?i)\bEXPLOIT_DEMO_ONLY\b|\bcurl\s+[^\n]*Authorization:"),
    ),
)


def _extract_docx(path: Path) -> str:
    document = Document(str(path))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    chunks.extend(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return "\n".join(chunks)


def _extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        values = (
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        return "\n".join(values)
    finally:
        workbook.close()


def _extract_pdf(path: Path) -> str:
    document = fitz.open(path)
    try:
        return "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()


class ResidualScanner:
    def __init__(self, key_store: MasterKeyStore) -> None:
        self._fingerprints = KeyedFingerprinter(key_store)

    def scan(
        self,
        path: Path,
        *,
        forbidden_literals: tuple[str, ...] = (),
    ) -> ResidualScanResult:
        output_bytes = path.read_bytes()
        output_sha256 = hashlib.sha256(output_bytes).hexdigest()
        try:
            text = {
                ".docx": _extract_docx,
                ".xlsx": _extract_xlsx,
                ".pdf": _extract_pdf,
            }[path.suffix.lower()](path)
        except (KeyError, OSError, ValueError, RuntimeError, fitz.FileDataError):
            finding = ResidualFinding(
                pattern_id="residual_scan_unavailable",
                category=ContentCategory.UNCLASSIFIED,
                location="output",
                value_hmac=self._fingerprints.digest(
                    "residual-scan-failure",
                    output_sha256,
                ),
                waivable=False,
            )
            return ResidualScanResult(False, False, output_sha256, (finding,))

        findings: list[ResidualFinding] = []
        for pattern in _PATTERNS:
            for match in pattern.expression.finditer(text):
                findings.append(
                    ResidualFinding(
                        pattern_id=pattern.id,
                        category=pattern.category,
                        location=f"text-offset:{match.start()}",
                        value_hmac=self._fingerprints.digest(
                            f"residual:{pattern.id}",
                            match.group(0),
                            scope=output_sha256,
                        ),
                        waivable=False,
                    )
                )
        for index, literal in enumerate(forbidden_literals, start=1):
            if literal and literal in text:
                findings.append(
                    ResidualFinding(
                        pattern_id=f"forbidden_literal_{index}",
                        category=ContentCategory.UNCLASSIFIED,
                        location="output",
                        value_hmac=self._fingerprints.digest(
                            "residual:forbidden-literal",
                            literal,
                            scope=output_sha256,
                        ),
                        waivable=False,
                    )
                )
        frozen = tuple(findings)
        return ResidualScanResult(
            safe=not frozen,
            waivable=all(finding.waivable for finding in frozen),
            output_sha256=output_sha256,
            findings=frozen,
        )


def residual_output_error(
    finding: ResidualFinding,
    file_kind: FileKind,
) -> ResidualOutputError:
    if finding.category is ContentCategory.CREDENTIAL_SECRET:
        rule_id = "PS-RESIDUAL-CREDENTIAL"
        location = "output:residual-scan:credential"
        evidence_ids = (
            ("EV-DOCX-003", "EV-GEN-004") if file_kind is FileKind.DOCX else ("EV-GEN-004",)
        )
    else:
        slug = finding.category.value.replace("_", "-").upper()
        rule_id = f"PS-RESIDUAL-{slug}"
        location = f"output:residual-scan:{finding.category.value}"
        evidence_ids = ("EV-GEN-002", "EV-GEN-004")
    return ResidualOutputError(
        "generated output contains a non-waivable residual risk",
        category=finding.category,
        location=location,
        rule_id=rule_id,
        evidence_ids=evidence_ids,
    )
