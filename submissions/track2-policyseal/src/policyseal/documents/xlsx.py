from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import MergedCell  # type: ignore[import-untyped]

from policyseal.documents.base import (
    CapabilityFinding,
    OutputValidation,
    ParsedDocument,
    PreflightDisposition,
    PreflightResult,
    PreviewArtifact,
    RenderedDocument,
)
from policyseal.documents.preflight import PreflightService
from policyseal.documents.residual import ResidualScanner, residual_output_error
from policyseal.documents.xlsx_mapping import (
    discover_workbook_mapping,
    metadata_category,
)
from policyseal.domain.enums import Action, FileKind
from policyseal.domain.errors import UnsafeOutputError, UnsupportedDocumentError
from policyseal.domain.models import DocumentRegion, PolicyDecision, SourceFingerprint

_DISPOSITION_RANK = {
    PreflightDisposition.PASS: 0,
    PreflightDisposition.REQUIRE_CONFIRMATION: 1,
    PreflightDisposition.REQUIRE_APPROVAL: 2,
    PreflightDisposition.BLOCK: 3,
    PreflightDisposition.PROCESSING_FAILED: 4,
}


def _stable_region_id(
    source_sha256: str,
    sheet_name: str,
    coordinate: str,
    normalized_header: str,
) -> str:
    payload = "\0".join((source_sha256, sheet_name, coordinate, normalized_header)).encode()
    return f"xlsx-{hashlib.sha256(payload).hexdigest()[:24]}"


def _location_parts(location: str) -> tuple[str, str]:
    prefix = "xlsx:"
    if not location.startswith(prefix):
        raise UnsafeOutputError("decision location is not an XLSX cell")
    parts = location.removeprefix(prefix).rsplit(":", 1)
    if len(parts) != 2 or not all(parts):
        raise UnsafeOutputError("decision location is malformed")
    return parts[0], parts[1]


def _decision_digest(decisions: tuple[PolicyDecision, ...]) -> str:
    payload = [
        {
            "action": item.action.value,
            "region_id": item.region_id,
            "rule_ids": item.rule_ids,
            "transform": item.transform,
        }
        for item in decisions
    ]
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()[:16]


def _status_value(value: object) -> str:
    normalized = str(value).strip().casefold()
    if "reject" in normalized:
        return "Rejected"
    if "approv" in normalized or "shortlist" in normalized:
        return "Approved"
    return "Pending"


def _shortlist_value(value: object) -> str:
    normalized = str(value).strip().casefold()
    if "not shortlist" in normalized or "reject" in normalized:
        return "Not shortlisted"
    if "shortlist" in normalized or "approv" in normalized:
        return "Shortlisted"
    return "Pending"


class XlsxAdapter:
    file_kind = FileKind.XLSX

    def __init__(self, preflight: PreflightService, residual_scanner: ResidualScanner) -> None:
        self._preflight = preflight
        self._residual = residual_scanner

    def preflight(self, path: Path) -> PreflightResult:
        base = self._preflight.inspect(path)
        if base.blocked:
            return base
        findings = list(base.findings)
        try:
            workbook = load_workbook(path, data_only=False, keep_links=False)
            try:
                mapping = discover_workbook_mapping(workbook)
            finally:
                workbook.close()
        except Exception:
            return PreflightResult(
                file_kind=FileKind.XLSX,
                disposition=PreflightDisposition.PROCESSING_FAILED,
                findings=(
                    CapabilityFinding(
                        code="xlsx_mapping_failed",
                        disposition=PreflightDisposition.PROCESSING_FAILED,
                        location="workbook",
                        message="Workbook columns cannot be mapped safely.",
                    ),
                ),
            )
        if not mapping.sheets:
            findings.append(
                CapabilityFinding(
                    code="xlsx_header_not_found",
                    disposition=PreflightDisposition.REQUIRE_CONFIRMATION,
                    location="workbook",
                    message="No controlled header row was found.",
                )
            )
        for sheet_name, column in mapping.unknown_columns:
            findings.append(
                CapabilityFinding(
                    code="xlsx_unknown_column",
                    disposition=PreflightDisposition.REQUIRE_CONFIRMATION,
                    location=f"xlsx:{sheet_name}:column:{column.index}",
                    message="An unknown column requires controlled mapping.",
                )
            )
        disposition = max(
            (finding.disposition for finding in findings),
            key=_DISPOSITION_RANK.__getitem__,
            default=PreflightDisposition.PASS,
        )
        return PreflightResult(FileKind.XLSX, disposition, tuple(findings))

    def parse(self, path: Path) -> ParsedDocument:
        preflight = self.preflight(path)
        if preflight.blocked:
            raise UnsupportedDocumentError("workbook failed the P0 capability preflight")
        source = SourceFingerprint.from_bytes(path.read_bytes())
        workbook = load_workbook(path, data_only=False, keep_links=False)
        regions: list[DocumentRegion] = []
        try:
            mapping = discover_workbook_mapping(workbook)
            for sheet_mapping in mapping.sheets:
                sheet = workbook[sheet_mapping.sheet_name]
                if sheet_mapping.sheet_name == "Supplier Comparison":
                    regions.extend(
                        self._metadata_regions(source.sha256, sheet, sheet_mapping.header_row)
                    )
                for row_index in range(sheet_mapping.header_row + 1, sheet.max_row + 1):
                    owner_id = None
                    if sheet_mapping.owner_column is not None:
                        owner_value = sheet.cell(row_index, sheet_mapping.owner_column).value
                        owner_id = str(owner_value).strip() if owner_value is not None else None
                    for column in sheet_mapping.columns:
                        cell = sheet.cell(row_index, column.index)
                        if cell.value is None:
                            continue
                        regions.append(
                            DocumentRegion(
                                id=_stable_region_id(
                                    source.sha256,
                                    sheet.title,
                                    cell.coordinate,
                                    column.normalized_header,
                                ),
                                file_kind=FileKind.XLSX,
                                location=f"xlsx:{sheet.title}:{cell.coordinate}",
                                order=len(regions),
                                text=str(cell.value),
                                category=column.category,
                                owner_id=owner_id,
                            )
                        )
        finally:
            workbook.close()
        if len({region.id for region in regions}) != len(regions):
            raise UnsupportedDocumentError("workbook region IDs are not unique")
        return ParsedDocument(
            source_path=path,
            source=source,
            file_kind=FileKind.XLSX,
            regions=tuple(regions),
            preflight=preflight,
        )

    def _metadata_regions(
        self,
        source_sha256: str,
        sheet: Any,
        header_row: int,
    ) -> list[DocumentRegion]:
        regions: list[DocumentRegion] = []
        for row_index in range(1, header_row):
            for column_index in range(1, sheet.max_column + 1):
                cell = sheet.cell(row_index, column_index)
                if cell.value is None:
                    continue
                category = metadata_category(sheet, row_index, column_index)
                regions.append(
                    DocumentRegion(
                        id=_stable_region_id(
                            source_sha256,
                            sheet.title,
                            cell.coordinate,
                            f"metadata:{category.value}",
                        ),
                        file_kind=FileKind.XLSX,
                        location=f"xlsx:{sheet.title}:{cell.coordinate}",
                        order=len(regions),
                        text=str(cell.value),
                        category=category,
                    )
                )
        return regions

    def render(
        self,
        parsed: ParsedDocument,
        decisions: tuple[PolicyDecision, ...],
        output_dir: Path,
    ) -> RenderedDocument:
        if parsed.file_kind is not FileKind.XLSX:
            raise UnsupportedDocumentError("XLSX adapter received a different file kind")
        if any(decision.action is Action.BLOCK for decision in decisions):
            raise UnsupportedDocumentError("blocked disclosure contexts cannot render an output")
        by_region = {decision.region_id: decision for decision in decisions}
        if len(by_region) != len(decisions) or set(by_region) != {
            region.id for region in parsed.regions
        }:
            raise UnsafeOutputError("every stable workbook region requires exactly one decision")

        workbook = load_workbook(parsed.source_path, data_only=False, keep_links=False)
        mapping = discover_workbook_mapping(workbook)
        rows_to_delete: dict[str, set[int]] = {}
        try:
            for region in parsed.regions:
                decision = by_region[region.id]
                sheet_name, coordinate = _location_parts(region.location)
                sheet = workbook[sheet_name]
                cell = sheet[coordinate]
                sheet_mapping = mapping.sheet(sheet_name)
                column = sheet_mapping.column(cell.column) if sheet_mapping is not None else None
                header = column.normalized_header if column is not None else ""
                if decision.action is Action.KEEP:
                    continue
                if decision.action is Action.REMOVE:
                    cell.value = None
                elif decision.action is Action.MASK:
                    self._mask(cell, decision)
                elif decision.action is Action.FILTER:
                    if decision.transform == "filter_other_owners":
                        rows_to_delete.setdefault(sheet_name, set()).add(cell.row)
                    else:
                        self._filter(cell, header, decision)
                elif decision.action is Action.GENERALIZE:
                    self._generalize(cell, header, decision)
                elif decision.action is Action.REQUIRE_APPROVAL:
                    cell.value = None
                else:
                    raise UnsafeOutputError("unsupported workbook action")

            for sheet_name, row_numbers in rows_to_delete.items():
                sheet = workbook[sheet_name]
                for row_number in sorted(row_numbers, reverse=True):
                    sheet.delete_rows(row_number, 1)

            external = any(
                rule_id.startswith(("PS-XLSX-AUDIT", "PS-XLSX-BIDDER"))
                for decision in decisions
                for rule_id in decision.rule_ids
            )
            if external and "Field Guide" in workbook.sheetnames:
                del workbook["Field Guide"]
            workbook.properties.creator = "PolicySeal"
            workbook.properties.lastModifiedBy = "PolicySeal"
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if not isinstance(cell, MergedCell):
                            cell.comment = None
            output_dir.mkdir(parents=True, exist_ok=True)
            output_path = output_dir / (
                f"policyseal-{parsed.source.sha256[:12]}-{_decision_digest(decisions)}.xlsx"
            )
            workbook.save(output_path)
        finally:
            workbook.close()

        validation = self.reopen_validate(output_path)
        residual = self._residual.scan(output_path)
        if not residual.safe:
            output_path.unlink(missing_ok=True)
            raise residual_output_error(residual.findings[0], FileKind.XLSX)
        if not validation.valid:
            output_path.unlink(missing_ok=True)
            raise UnsafeOutputError("generated workbook failed structure or residual validation")
        return RenderedDocument(
            output_path=output_path,
            decisions=decisions,
            validation=validation,
        )

    @staticmethod
    def _mask(cell: Any, decision: PolicyDecision) -> None:
        if decision.transform != "mask_last_four":
            raise UnsafeOutputError("unsupported workbook mask")
        digits = "".join(character for character in str(cell.value) if character.isdigit())
        if len(digits) < 4:
            raise UnsafeOutputError("payment account cannot be masked deterministically")
        cell.value = f"Account ending {digits[-4:]}"

    @staticmethod
    def _filter(cell: Any, header: str, decision: PolicyDecision) -> None:
        if decision.transform == "keep_name_work_email":
            if header not in {"contact name", "work email"}:
                cell.value = None
            return
        if decision.transform == "keep_audit_amount":
            if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
                cell.value = None
            return
        raise UnsafeOutputError("unsupported workbook filter")

    @staticmethod
    def _generalize(cell: Any, header: str, decision: PolicyDecision) -> None:
        if decision.transform == "generalize_approval_status":
            cell.value = _status_value(cell.value)
            return
        if decision.transform == "generalize_shortlist_status":
            cell.value = _shortlist_value(cell.value) if header == "decision status" else None
            return
        raise UnsafeOutputError("unsupported workbook generalization")

    def reopen_validate(self, path: Path) -> OutputValidation:
        try:
            workbook = load_workbook(path, data_only=False, keep_links=False)
            try:
                has_primary = "Supplier Comparison" in workbook.sheetnames
                has_formula = any(
                    isinstance(cell.value, str) and cell.value.startswith("=")
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                )
            finally:
                workbook.close()
        except Exception:
            return OutputValidation(False, False, ("xlsx_reopen_failed",))
        findings = tuple(
            code
            for condition, code in (
                (not has_primary, "xlsx_primary_sheet_missing"),
                (has_formula, "xlsx_formula_in_output"),
            )
            if condition
        )
        return OutputValidation(not findings, has_primary, findings)

    def preview(self, path: Path, output_dir: Path) -> tuple[PreviewArtifact, ...]:
        del path, output_dir
        return ()
