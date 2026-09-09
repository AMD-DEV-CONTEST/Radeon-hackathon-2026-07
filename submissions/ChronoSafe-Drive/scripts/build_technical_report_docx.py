"""Build the reviewer-ready English technical report as an editable DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BRAND_RED = "9F1239"
NAVY = "1F3A5F"
INK = "20262E"
MUTED = "64748B"
LIGHT_FILL = "F2F4F7"
GRID = "94A3B8"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=PROJECT_ROOT / "docs" / "TECHNICAL_REPORT.md",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "docs" / "TECHNICAL_REPORT.docx",
    )
    return parser.parse_args()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color: str = GRID, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def apply_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_table_borders(table)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=8, color=MUTED)


def set_paragraph_border(paragraph, *, side: str, color: str, size: int = 18, space: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_tokens = {
        "Heading 1": (16, BRAND_RED, 16, 8),
        "Heading 2": (13, NAVY, 12, 6),
        "Heading 3": (12, "374151", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for section in document.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        run = hp.add_run("ChronoSafe-Drive | AMD ROCm Physical AI")
        set_run_font(run, size=8, color=MUTED)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_after = Pt(0)
        label = fp.add_run("Page ")
        set_run_font(label, size=8, color=MUTED)
        add_page_field(fp)

    document.core_properties.title = "ChronoSafe-Drive Technical Report"
    document.core_properties.subject = "AMD Hackathon 2026 - Track 3: Physical AI"
    document.core_properties.author = "Wang Zhenyu; Wang Guanyu; Han Liuxuan"
    document.core_properties.keywords = "AMD Radeon, ROCm, Physical AI, autonomous driving"


def next_numbering_id(numbering, tag: str) -> int:
    values = [int(node.get(qn(f"w:{tag}"))) for node in numbering.findall(qn(f"w:{tag[:-2]}")) if node.get(qn(f"w:{tag}"))]
    return max(values, default=0) + 1


def create_numbering(document: DocumentType, *, ordered: bool) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(level_text)
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "left")
    level.append(justify)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, default_size: float = 11, default_color: str = INK) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=default_color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(default_size - 1, 7.5), color=BRAND_RED)
        else:
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                run = paragraph.add_run(link_match.group(1))
                set_run_font(run, size=default_size, color="0563C1")
                run.underline = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=default_size, color=default_color)


def add_body_paragraph(document: DocumentType, text: str):
    paragraph = document.add_paragraph()
    add_inline(paragraph, text)
    paragraph.paragraph_format.widow_control = True
    return paragraph


def add_quote(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color="374151", italic=True)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFF7ED")
    p_pr.append(shading)
    set_paragraph_border(paragraph, side="left", color="DC2626", size=20, space=7)


def add_code_block(document: DocumentType, lines: Iterable[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, name="Consolas", size=8.5, color="374151")
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F8FAFC")
    p_pr.append(shading)
    for side in ("top", "left", "bottom", "right"):
        set_paragraph_border(paragraph, side=side, color="CBD5E1", size=6, space=4)


def choose_table_widths(rows: list[list[str]]) -> list[int]:
    columns = max(len(row) for row in rows)
    if columns == 2:
        return [4300, 5060]
    if columns == 3:
        return [2200, 2500, 4660]
    if columns == 4:
        return [2500, 1900, 2960, 2000]
    if columns == 7:
        return [1200, 700, 1100, 1100, 1100, 2100, 2060]
    base = CONTENT_WIDTH_DXA // columns
    widths = [base] * columns
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_markdown_table(document: DocumentType, lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    apply_table_geometry(table, choose_table_widths(rows))
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(
                paragraph,
                value,
                default_size=8.5,
                default_color="FFFFFF" if row_index == 0 else "475569",
            )
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in paragraph.runs:
                    run.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_architecture_flow(document: DocumentType) -> None:
    steps = [
        "Ingest MetaDrive physics simulation and nuScenes camera/LiDAR streams.",
        "Apply controlled single, compound, and timestamp degradation.",
        "Build an eight-frame window of 15-D causal sensor-health features.",
        "Run the shared GRU with camera, LiDAR, timing, and continuous-risk heads.",
        "Calibrate a worst-domain conformal upper-risk bound.",
        "Drive the NORMAL -> DEGRADED -> MINIMAL_RISK closed-loop controller.",
    ]
    num_id = create_numbering(document, ordered=True)
    for step in steps:
        paragraph = document.add_paragraph()
        apply_numbering(paragraph, num_id)
        add_inline(paragraph, step)


def add_cover(document: DocumentType) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("AMD HACKATHON 2026 | TRACK 3: PHYSICAL AI")
    set_run_font(run, size=10, color=BRAND_RED, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("ChronoSafe-Drive")
    set_run_font(run, size=28, color=BRAND_RED, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run(
        "Self-Diagnosis, Risk Calibration, and Minimal-Risk Maneuvers under "
        "Multi-Sensor Perception Degradation"
    )
    set_run_font(run, size=14, color="334155")

    metric_rows = [
        ["Closed-loop MRM", "Timestamp robustness", "AMD ROCm inference"],
        ["20/20 success\n0% premature", "95.08%-99.34%\nstress accuracy", "0.559 ms p50\n1.37M seq/s"],
    ]
    table = document.add_table(rows=2, cols=3)
    apply_table_geometry(table, [3120, 3120, 3120])
    for row_index, row in enumerate(metric_rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(
                run,
                size=9 if row_index == 0 else 11,
                color="FFFFFF" if row_index == 0 else "111827",
                bold=True,
            )
            set_cell_shading(cell, NAVY if row_index == 0 else LIGHT_FILL)

    details = document.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.paragraph_format.space_before = Pt(28)
    details.paragraph_format.space_after = Pt(4)
    run = details.add_run("AMD Radeon gfx1100 | ROCm/HIP 7.2 | PyTorch 2.9.1 ROCm")
    set_run_font(run, size=10, color="334155")
    details2 = document.add_paragraph()
    details2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details2.paragraph_format.space_after = Pt(26)
    run = details2.add_run("nuScenes v1.0-trainval | 20 unseen MetaDrive closed-loop seeds")
    set_run_font(run, size=10, color="334155")
    team = document.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team.paragraph_format.space_after = Pt(4)
    run = team.add_run("Wang Zhenyu | Wang Guanyu | Han Liuxuan")
    set_run_font(run, size=10, color="334155", bold=True)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date.add_run("Team submission | August 2026")
    set_run_font(run, size=11, color=MUTED, italic=True)
    document.add_page_break()


def build_markdown_body(document: DocumentType, text: str) -> None:
    lines = text.splitlines()
    index = 0
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            add_body_paragraph(document, " ".join(paragraph_lines))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if index < 4 and stripped.startswith("#"):
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            language = stripped.removeprefix("```").strip()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            if language == "mermaid":
                add_architecture_flow(document)
            else:
                add_code_block(document, block)
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            source_level = len(heading.group(1))
            word_level = 1 if source_level == 2 else 2 if source_level == 3 else 3
            if heading.group(2) == "Abstract":
                word_level = 1
            paragraph = document.add_paragraph(style=f"Heading {word_level}")
            add_inline(paragraph, heading.group(2), default_size={1: 16, 2: 13, 3: 12}[word_level])
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(
                    BRAND_RED if word_level == 1 else NAVY if word_level == 2 else "374151"
                )
            index += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().removeprefix(">").strip())
                index += 1
            add_quote(document, " ".join(quote_lines))
            continue
        if re.match(r"^[-*]\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            ordered = bool(re.match(r"^\d+\.\s+", stripped))
            item_texts: list[str] = []
            while index < len(lines):
                raw = lines[index]
                item = raw.strip()
                match = re.match(r"^\d+\.\s+(.+)$", item) if ordered else re.match(r"^[-*]\s+(.+)$", item)
                if match:
                    item_texts.append(match.group(1))
                    index += 1
                    continue
                if item_texts and raw[:1].isspace() and item:
                    item_texts[-1] = f"{item_texts[-1]} {item}"
                    index += 1
                    continue
                break
            num_id = create_numbering(document, ordered=ordered)
            for item in item_texts:
                paragraph = document.add_paragraph()
                apply_numbering(paragraph, num_id)
                add_inline(paragraph, item)
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def append_benchmark_figure(document: DocumentType) -> None:
    benchmark = PROJECT_ROOT / "evidence" / "rocm_inference_benchmark_seed2026" / "benchmark.png"
    if not benchmark.is_file():
        return
    document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("Appendix A. AMD ROCm Benchmark Figure")
    caption = document.add_paragraph()
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_after = Pt(6)
    run = caption.add_run(
        "Figure 1. Device-resident inference latency and throughput on the "
        "AMD Radeon gfx1100 GPU."
    )
    set_run_font(run, size=9, color=MUTED, italic=True)
    document.add_picture(str(benchmark), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    args = parse_args()
    source = args.input if args.input.is_absolute() else PROJECT_ROOT / args.input
    output = args.output if args.output.is_absolute() else PROJECT_ROOT / args.output
    if not source.is_file():
        raise FileNotFoundError(source)
    document = Document()
    configure_document(document)
    add_cover(document)
    build_markdown_body(document, source.read_text(encoding="utf-8"))
    append_benchmark_figure(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"TECHNICAL REPORT DOCX PASS: {output}")


if __name__ == "__main__":
    main()
