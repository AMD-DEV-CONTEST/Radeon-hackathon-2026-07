"""DOCX paragraph and table parser."""

from pathlib import Path

from docx import Document

from buildguard.rag.parsers.base import (
    DocumentParser,
    extract_article,
    extract_chapter,
    extract_effective_date,
    file_checksum,
)
from buildguard.rag.schemas import ParsedDocument, ParsedSection


class DocxParser(DocumentParser):
    extensions = frozenset({".docx"})

    def parse(self, path: Path) -> ParsedDocument:
        source = self.validate_path(path)
        checksum = file_checksum(source)
        document = Document(str(source))
        all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        effective_date = extract_effective_date(all_text)
        sections: list[ParsedSection] = []
        chapter: str | None = None
        title: str | None = None
        for paragraph in document.paragraphs:
            content = paragraph.text.strip()
            if not content:
                continue
            style = paragraph.style.name.lower() if paragraph.style else ""
            if "heading" in style or "标题" in style:
                title = content
            chapter = extract_chapter(content) or chapter
            sections.append(
                ParsedSection(
                    document_name=source.name,
                    source_path=source,
                    chapter=chapter,
                    article=extract_article(content),
                    title=title,
                    content=content,
                    effective_date=effective_date,
                    checksum=checksum,
                )
            )
        for table in document.tables:
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
            if not rows:
                continue
            sections.append(
                ParsedSection(
                    document_name=source.name,
                    source_path=source,
                    chapter=chapter,
                    title=title,
                    content="\n".join(rows),
                    effective_date=effective_date,
                    checksum=checksum,
                    is_table=True,
                    table_header=rows[0],
                )
            )
        return ParsedDocument(
            document_name=source.name,
            source_path=source,
            checksum=checksum,
            sections=sections,
            warnings=[] if sections else ["DOCX document contains no text or tables"],
        )
