"""
Insert captured screenshots into the PPTX, replacing placeholder shapes.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

PPTX_PATH = "/workspace/persistence/AMD-hacthon/amd-inventory-agent V1.1/docs/demo_video_materials/InventoryDoc_Agent_Demo.pptx"
ASSETS_DIR = "/workspace/persistence/AMD-hacthon/amd-inventory-agent V1.1/docs/demo_video_materials/slides/demo_video_assets"
OUTPUT_PATH = "/workspace/persistence/AMD-hacthon/amd-inventory-agent V1.1/docs/demo_video_materials/InventoryDoc_Agent_Demo.pptx"

prs = Presentation(PPTX_PATH)

# Mapping: slide_index (0-based) -> list of (img_path, target_shape_name_fragment)
slide_images = {
    0: [  # 封面
        (os.path.join(ASSETS_DIR, "slide_command_center.png"), "Slide 1"),
    ],
    1: [  # 真实业务问题
        (os.path.join(ASSETS_DIR, "crop_core_metrics.png"), "Slide 2"),
    ],
    2: [  # Agent做了什么
        (os.path.join(ASSETS_DIR, "slide_about_agent.png"), "Slide 3"),
    ],
    3: [  # 风险可视化
        (os.path.join(ASSETS_DIR, "slide_risk_visualization.png"), "Slide 4"),
    ],
    4: [  # 审计追踪
        (os.path.join(ASSETS_DIR, "slide_audit_trace.png"), "Slide 5"),
    ],
    5: [  # 管理层报告
        (os.path.join(ASSETS_DIR, "slide_management_report.png"), "Slide 6"),
    ],
    6: [  # 追问
        (os.path.join(ASSETS_DIR, "slide_followup_qa.png"), "Slide 7"),
    ],
    7: [  # Track2对齐
        (os.path.join(ASSETS_DIR, "slide_about_agent.png"), "Slide 8"),
    ],
    8: [  # 总结
        (os.path.join(ASSETS_DIR, "slide_command_center.png"), "Slide 9"),
    ],
}

for slide_idx, images in slide_images.items():
    slide = prs.slides[slide_idx]
    for img_path, label_fragment in images:
        if not os.path.exists(img_path):
            print(f"  [SKIP] Image not found: {img_path}")
            continue

        # Find the placeholder shape by checking text content
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text
                if label_fragment in text:
                    left = shape.left
                    top = shape.top
                    width = shape.width
                    height = shape.height
                    z_order = 0  # We'll track insertion order

                    # Delete the placeholder shape
                    sp = shape._element
                    sp.getparent().remove(sp)
                    print(f"  Slide {slide_idx+1}: removed placeholder, inserting {os.path.basename(img_path)}")
                    break
        else:
            # No placeholder found; place at default position
            print(f"  Slide {slide_idx+1}: no placeholder found for '{label_fragment}', using default position")
            left = Inches(0.6)
            top = Inches(1.5)
            width = Inches(12)
            height = Inches(5)

        # Insert the image
        slide.shapes.add_picture(img_path, left, top, width, height)
        print(f"  Slide {slide_idx+1}: inserted {os.path.basename(img_path)}")

prs.save(OUTPUT_PATH)
print(f"\n✅ PPTX saved: {OUTPUT_PATH}")
