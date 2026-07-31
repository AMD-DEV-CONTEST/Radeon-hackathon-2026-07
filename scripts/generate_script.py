"""
InventoryDoc Agent — PPT 转脚本工具 (Step 1)

读取 PPTX 每页标题和正文，生成 narration.md 配音脚本。
"""

import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    print("错误: 请先安装 python-pptx: pip install python-pptx")
    sys.exit(1)


def extract_slides(pptx_path: str) -> list[dict]:
    prs = Presentation(pptx_path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        body_lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                # 第一个非空文本框作为标题
                if not title:
                    # 判断是否为标题形状
                    if shape.is_placeholder and shape.placeholder_format.idx == 0:
                        title = text
                    else:
                        title = text
                else:
                    body_lines.append(text)
        slides.append({"number": i, "title": title, "body": body_lines})
    return slides


def generate_script(slides: list[dict]) -> str:
    lines = []
    for s in slides:
        lines.append(f"## Slide {s['number']:02d}: {s['title']}")
        lines.append("")
        # 根据标题自动生成旁白
        voiceover = _auto_narration(s)
        lines.append(voiceover)
        lines.append("")
        lines.append("---")
        lines.append("")
    return "\n".join(lines)


def _auto_narration(s: dict) -> str:
    """根据幻灯片内容自动生成旁白草稿"""
    title = s["title"]
    body = s["body"]

    if s["number"] == 1:
        return (
            "欢迎观看 InventoryDoc Agent 的产品演示。"
            "这是一个基于 AMD Radeon GPU 和 ROCm 的本地私有 AI 智能体，"
            "专为库存盘点审计场景打造。"
        )

    if "痛点" in title:
        return (
            "仓储盘点面临四大痛点。"
            "第一，数据分散在十多个文件中，编码不统一。"
            "第二，噪音干扰严重，NOISE 数据和异常记录混杂其中。"
            "第三，人工分析需要两到三天，效率太低。"
            "第四，容易出错，基准选错、成本用了售价、数字前后不一致。"
        )

    if "方案" in title:
        return (
            "我们的方案是 InventoryDoc Agent。"
            "你只需要上传十份盘点数据，它就能在十分钟内自动生成可直接上会的调账材料。"
            "它具备本地知识检索、多步骤任务规划、本地多轮记忆和工具调用等核心能力，"
            "数据全程不出设备，零云端依赖。"
        )

    if "工作流" in title or "10 步" in title:
        return (
            "这是 InventoryDoc Agent 的十步确定性工作流。"
            "从多格式文件加载开始，经过智能噪音识别、基准选择三维论证、"
            "扫描日志清洗、盘点结果加载、差异金额计算、日志关联三分类、"
            "库位状态校验、品类维度分析，到最后交付物生成，"
            "每一步都确保数据准确性和可追溯性。"
        )

    if "噪音" in title:
        return (
            "第一个创新点是噪音自动识别。"
            "Agent 能自动识别三类噪音："
            "RF 扫码环境噪声，自动过滤六十八条干扰数据；"
            "异常记录，历史库存中超过两百条异常自动标记；"
            "以及不适用文件，如粒度不匹配的月报自动排除。"
            "人工做这一步需要逐行检查，Agent 三秒完成。"
        )

    if "基准" in title:
        return (
            "第二个创新点是基准选择三维论证。"
            "Agent 从编码体系、时间窗口和数据粒度三个维度自动对比，"
            "选对基准。"
            "选错基准，整个分析就白做了。"
        )

    if "日志" in title or "三分类" in title:
        return (
            "第三个创新点是日志关联三分类。"
            "Agent 按净变动量进行精准比对，"
            "将一百五十九个差异 SKU 分为三类："
            "十四个完全可解释，有对应的出入库记录；"
            "二十一个部分可解释，有操作但数量不够；"
            "一百二十四个没有日志关联，这些是需要去仓库现场实物复核的重点。"
        )

    if "结果" in title or "演示" in title:
        return (
            "这是演示结果。"
            "系统库存总计两万一千一百四十四件，净差异负九百七十一件，"
            "总体差异率百分之四点六，"
            "净差异金额负九十八万九千九百元。"
            "质量校验七项全部通过。"
        )

    if "ROCm" in title or "AMD" in title:
        return (
            "为什么选择 AMD ROCm？"
            "因为完全开源，没有供应商锁定；"
            "硬件成本更低，企业部署更友好；"
            "vLLM 原生支持，高性能推理；"
            "数据安全，本地部署，数据不出设备；"
            "零 API 费用，一次部署，持续使用。"
        )

    if "价值" in title or "谢谢" in title:
        return (
            "业务价值方面，分析耗时从两到三天缩短到十分钟，提升百分之九十九。"
            "感谢您的关注，欢迎访问我们的 GitHub 仓库了解更多详情。"
        )

    # 兜底
    parts = [f"接下来看第{s['number']}页：{title}。"]
    if body:
        first_bullet = body[0][:60]
        parts.append(f"主要内容包括：{first_bullet}。")
    return " ".join(parts)


def main():
    import argparse

    parser = argparse.ArgumentParser(description="从 PPTX 生成配音脚本")
    parser.add_argument("--input", required=True, help="输入的 PPTX 文件路径")
    parser.add_argument("--output", default="narration.md", help="输出的脚本文件路径")
    args = parser.parse_args()

    pptx_path = Path(args.input)
    if not pptx_path.exists():
        print(f"错误: 找不到文件 {pptx_path}")
        sys.exit(1)

    print(f"读取 PPTX: {pptx_path}")
    slides = extract_slides(str(pptx_path))
    print(f"共发现 {len(slides)} 页幻灯片")

    script = generate_script(slides)
    output_path = Path(args.output)
    output_path.write_text(script, encoding="utf-8")
    print(f"脚本已生成: {output_path}")
    print(f"请编辑 {output_path} 微调旁白内容后继续下一步。")


if __name__ == "__main__":
    main()