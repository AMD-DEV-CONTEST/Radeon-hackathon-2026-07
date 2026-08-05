"""
Generate 9-slide PPTX skeleton for InventoryDoc Agent demo video.
Layout + text + screenshot placeholder areas.
User manually inserts screenshots after generation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Color palette ──
BG_DARK = RGBColor(0x0D, 0x11, 0x17)
AMD_RED = RGBColor(0xFF, 0x00, 0x00)
TECH_BLUE = RGBColor(0x58, 0xA6, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x8B, 0x94, 0x9E)
DARK_CARD = RGBColor(0x16, 0x1B, 0x22)
PLACEHOLDER_BORDER = RGBColor(0x30, 0x3D, 0x4D)

# ── Slide dimensions (16:9) ──
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# Use blank layout
blank_layout = prs.slide_layouts[6]


def set_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_multiline_textbox(slide, left, top, width, height, lines,
                          font_size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
                          line_spacing=1.5, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, (text, fs, c, b) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(fs)
        p.font.color.rgb = c
        p.font.bold = b
        p.font.name = font_name
        p.alignment = align
        p.space_after = Pt(4)
    return txBox


def add_image_placeholder(slide, left, top, width, height, label="截图占位区"):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_CARD
    shape.line.color.rgb = PLACEHOLDER_BORDER
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"[ {label} ]"
    p.font.size = Pt(14)
    p.font.color.rgb = GRAY
    p.font.name = "Microsoft YaHei"
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape


def add_amd_red_line(slide, left, top, width):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, Pt(3)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = AMD_RED
    line.line.fill.background()
    return line


# ══════════════════════════════════════════════════════════════
# SLIDE 1 — 封面
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

# Screenshot placeholder (main visual, ~70%)
add_image_placeholder(slide, Inches(0.6), Inches(0.6), Inches(12.1), Inches(5.0),
                      "Slide 1: 指挥舱长截图 command_center.png")

# Top-left title overlaid
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(8), Inches(1.0),
            "InventoryDoc Agent", font_size=36, color=WHITE, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.7), Inches(8), Inches(0.6),
            "本地私有库存差异审计Agent", font_size=20, color=GRAY)

# Bottom tagline
add_amd_red_line(slide, Inches(0.6), Inches(5.9), Inches(2.0))
add_textbox(slide, Inches(0.6), Inches(6.1), Inches(12), Inches(0.5),
            "10份业务文件 → 10步审计工作流 → 质量门禁 → 管理层报告",
            font_size=14, color=GRAY)

# Track 2 badge bottom-right
add_textbox(slide, Inches(9.5), Inches(6.5), Inches(3.5), Inches(0.5),
            "Track 2 · Private AI Agent", font_size=13, color=TECH_BLUE,
            align=PP_ALIGN.RIGHT)

# Title text at bottom
add_textbox(slide, Inches(0.6), Inches(6.6), Inches(8), Inches(0.5),
            "这不是概念图，是已经跑通的库存审计Agent",
            font_size=16, color=WHITE, bold=False)

# ══════════════════════════════════════════════════════════════
# SLIDE 2 — 真实业务问题
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "真实业务问题：10份文件，528个SKU",
            font_size=30, color=WHITE, bold=True)

add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Big numbers (left column)
add_multiline_textbox(slide, Inches(0.6), Inches(1.5), Inches(4.5), Inches(3.0), [
    ("10份数据", 40, AMD_RED, True),
    ("528个SKU", 40, WHITE, True),
    ("净差异 -971件", 36, WHITE, True),
    ("净差异金额 -989,900.97元", 36, WHITE, True),
])

# Subtitle
add_textbox(slide, Inches(0.6), Inches(4.8), Inches(5), Inches(0.5),
            "人工分析2-3天，最容易错在基准、成本和日志噪音",
            font_size=14, color=GRAY)

# Screenshot placeholder (right side, ~50%)
add_image_placeholder(slide, Inches(5.5), Inches(1.3), Inches(7.2), Inches(5.5),
                      "Slide 2: 核心数字卡片 crop_core_metrics.png\n或指挥舱截图+圈注")

# ══════════════════════════════════════════════════════════════
# SLIDE 3 — Agent做了什么
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "Agent做了什么：从数据到调账材料",
            font_size=30, color=WHITE, bold=True)

add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Left: Input-Process-Output
add_multiline_textbox(slide, Inches(0.6), Inches(1.5), Inches(5.0), Inches(3.5), [
    ("输入：10份业务文件", 22, TECH_BLUE, True),
    ("  库存快照 · 盘点表 · 扫描日志 · 主数据...", 16, GRAY, False),
    ("", 12, GRAY, False),
    ("过程：10步确定性工作流", 22, TECH_BLUE, True),
    ("  Load → Noise → Baseline → Clean → Count → Variance →", 16, GRAY, False),
    ("  Correlation → Location → Category → Deliverables", 16, GRAY, False),
    ("", 12, GRAY, False),
    ("输出：Excel + 报告 + JSON + 管理层摘要", 22, TECH_BLUE, True),
    ("  盘点差异明细.xlsx · 分析报告.md · 管理层报告.md", 16, GRAY, False),
])

# Right: screenshot placeholder
add_image_placeholder(slide, Inches(6.0), Inches(1.5), Inches(6.7), Inches(4.5),
                      "Slide 3: 关于页 about_agent.png")

# Bottom annotation
add_amd_red_line(slide, Inches(0.6), Inches(6.5), Inches(12))
add_textbox(slide, Inches(0.6), Inches(6.7), Inches(12), Inches(0.5),
            "事实计算 pandas  ·  报告表达 LLM  ·  数据边界本地",
            font_size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 4 — 风险可视化
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "风险不是文字描述，是可视化定位",
            font_size=30, color=WHITE, bold=True)
add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Screenshot (main area)
add_image_placeholder(slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.2),
                      "Slide 4: 风险可视化 risk_visualization.png")

# Tagline overlay
add_textbox(slide, Inches(0.8), Inches(6.8), Inches(12), Inches(0.5),
            "差异率4.6%  ·  异常品类（浴室家具16.0%） ·  Top风险SKU  ·  金额影响",
            font_size=16, color=TECH_BLUE, align=PP_ALIGN.CENTER)

# Workflow header
add_textbox(slide, Inches(0.6), Inches(1.4), Inches(3), Inches(0.5),
            '"哪里有差异" → "哪里先处理"',
            font_size=14, color=GRAY)

# ══════════════════════════════════════════════════════════════
# SLIDE 5 — 审计追踪
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "最关键：每个结论都有审计追踪",
            font_size=30, color=WHITE, bold=True)
add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Screenshot (main area)
add_image_placeholder(slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.8),
                      "Slide 5: 审计追踪 audit_trace.png\n（含质量门禁7/7、run_id）")

# Bottom key points
add_textbox(slide, Inches(0.6), Inches(6.4), Inches(12), Inches(0.5),
            "每一步都有证据  ·  每个数字可回溯",
            font_size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.6), Inches(6.9), Inches(12), Inches(0.5),
            "run_id  ·  输入文件记录  ·  步骤耗时  ·  质量门禁(7/7)  ·  输出材料",
            font_size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 6 — 管理层报告
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "LLM只写报告，不参与算账",
            font_size=30, color=WHITE, bold=True)
add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Screenshot (main area)
add_image_placeholder(slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.5),
                      "Slide 6: 管理层报告 management_report.png")

# Bottom
add_textbox(slide, Inches(0.6), Inches(6.1), Inches(12), Inches(0.5),
            "管理层报告  ·  LLM只负责表达，不负责事实计算",
            font_size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.6), Inches(6.6), Inches(12), Inches(0.5),
            "核心发现  ·  风险分级  ·  Top异常SKU  ·  处置建议  ·  人审边界",
            font_size=14, color=GRAY, align=PP_ALIGN.CENTER)

# Add report mode badge
add_textbox(slide, Inches(10.5), Inches(0.4), Inches(2.5), Inches(0.4),
            "生成模式: demo / mock / real", font_size=11, color=GRAY, align=PP_ALIGN.RIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 7 — 追问
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "管理层可以继续追问",
            font_size=30, color=WHITE, bold=True)
add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Screenshot (main area)
add_image_placeholder(slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.5),
                      "Slide 7: 追问 followup_qa.png")

add_textbox(slide, Inches(0.6), Inches(6.1), Inches(12), Inches(0.5),
            "不是静态报告  ·  而是可交互的审计助手",
            font_size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.6), Inches(6.6), Inches(12), Inches(0.5),
            "问差异  ·  问品类  ·  问Top SKU  ·  问证据来源",
            font_size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 8 — Track2 对齐
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
            "对齐Track2：本地部署 + Agent能力 + GPU报告",
            font_size=28, color=WHITE, bold=True)
add_amd_red_line(slide, Inches(0.6), Inches(1.1), Inches(3.0))

# Left: capabilities list
add_multiline_textbox(slide, Inches(0.6), Inches(1.5), Inches(5.5), Inches(4.0), [
    ("Track 2 · Private AI Agent", 22, TECH_BLUE, True),
    ("", 10, GRAY, False),
    ("✓  本地私有部署", 20, WHITE, False),
    ("    数据在本地处理，不上传云端", 14, GRAY, False),
    ("", 10, GRAY, False),
    ("✓  Multi-step Planning（10步工作流）", 20, WHITE, False),
    ("    完整任务拆解为可追踪步骤", 14, GRAY, False),
    ("", 10, GRAY, False),
    ("✓  Tool Calling（多格式文件读写）", 20, WHITE, False),
    ("    Excel / CSV / JSON / Markdown", 14, GRAY, False),
    ("", 10, GRAY, False),
    ("✓  Memory & Audit Trail（运行记忆）", 20, WHITE, False),
    ("    run_id · 中间状态 · 质量门禁", 14, GRAY, False),
    ("", 10, GRAY, False),
    ("✓  Quality Control（7条质量门禁）", 20, WHITE, False),
    ("    关键校验失败 → 报告不生成", 14, GRAY, False),
])

# Right: screenshot
add_image_placeholder(slide, Inches(6.5), Inches(1.5), Inches(6.2), Inches(4.5),
                      "Slide 8: 关于页 + 管理层报告")

# Bottom
add_amd_red_line(slide, Inches(0.6), Inches(6.5), Inches(12))
add_textbox(slide, Inches(0.6), Inches(6.7), Inches(12), Inches(0.5),
            "管理层报告可通过 AMD ROCm 本地LLM 推理生成",
            font_size=15, color=TECH_BLUE, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 9 — 总结
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide)

# Background screenshot (light)
add_image_placeholder(slide, Inches(0.0), Inches(0.0), SLIDE_W, SLIDE_H,
                      "Slide 9: 指挥舱截图（背景淡出）")

add_textbox(slide, Inches(0.6), Inches(0.5), Inches(12), Inches(0.7),
            "总结：分钟级、可追溯、可上会",
            font_size=30, color=WHITE, bold=True)
add_amd_red_line(slide, Inches(0.6), Inches(1.2), Inches(3.0))

# Center keywords
add_multiline_textbox(slide, Inches(1.0), Inches(2.0), Inches(11), Inches(3.0), [
    ("从2-3天手工分析", 28, GRAY, False),
    ("到分钟级审计流程", 28, WHITE, True),
], align=PP_ALIGN.CENTER)

# Three big keywords
add_multiline_textbox(slide, Inches(0.5), Inches(3.8), Inches(12.3), Inches(2.0), [
    ("真实业务     本地私有     审计可信", 32, AMD_RED, True),
], align=PP_ALIGN.CENTER)

# Bottom
add_textbox(slide, Inches(0.6), Inches(6.5), Inches(12), Inches(0.5),
            "InventoryDoc Agent  ·  Built for AMD Radeon GPU + ROCm",
            font_size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# ── Save ──
output_path = "/workspace/persistence/AMD-hacthon/amd-inventory-agent V1.1/docs/demo_video_materials/InventoryDoc_Agent_Demo.pptx"
prs.save(output_path)
print(f"PPTX saved: {output_path}")
print(f"Slides: {len(prs.slides)}")
