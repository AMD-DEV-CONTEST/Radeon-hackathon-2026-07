#!/usr/bin/env python3
"""Create PPTX presentation for InventoryDoc Agent."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# AMD Red color
AMD_RED = RGBColor(237, 28, 36)
DARK_GRAY = RGBColor(51, 51, 51)
LIGHT_GRAY = RGBColor(128, 128, 128)
WHITE = RGBColor(255, 255, 255)

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(240, 240, 240)
    
    title = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "InventoryDoc Agent"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    p.alignment = PP_ALIGN.CENTER
    
    subtitle = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(11), Inches(1))
    tf = subtitle.text_frame
    p = tf.paragraphs[0]
    p.text = "库存盘点智能审计系统"
    p.font.size = Pt(28)
    p.font.color.rgb = DARK_GRAY
    p.alignment = PP_ALIGN.CENTER
    
    desc = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11), Inches(1))
    tf = desc.text_frame
    p = tf.paragraphs[0]
    p.text = "基于 AMD Radeon GPU + ROCm 的本地私有 AI Agent"
    p.font.size = Pt(18)
    p.font.color.rgb = LIGHT_GRAY
    p.alignment = PP_ALIGN.CENTER
    
    track = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11), Inches(0.5))
    tf = track.text_frame
    p = tf.paragraphs[0]
    p.text = "赛道二：私有 AI Agent 开发与本地部署"
    p.font.size = Pt(16)
    p.font.color.rgb = AMD_RED
    p.alignment = PP_ALIGN.CENTER
    
    # Slide 2: Pain Points
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "仓储盘点痛点"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    pain_points = [
        ("数据分散", "10+ 份文件（CSV/XLSX/TXT），编码不统一"),
        ("噪音干扰", "NOISE 干扰数据、异常记录混杂"),
        ("分析耗时", "人工分析需 2-3 天"),
        ("容易出错", "基准选错、成本用售价、数字不一致"),
    ]
    
    for i, (title_text, desc_text) in enumerate(pain_points):
        y = Inches(1.5 + i * 1.3)
        box = slide.shapes.add_textbox(Inches(1), y, Inches(11), Inches(1))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = f"❌ {title_text}"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = DARK_GRAY
        
        p2 = tf.add_paragraph()
        p2.text = desc_text
        p2.font.size = Pt(18)
        p2.font.color.rgb = LIGHT_GRAY
    
    # Slide 3: Solution
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "我们的方案：InventoryDoc Agent"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    solution = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = solution.text_frame
    tf.word_wrap = True
    
    lines = [
        "上传 10 份盘点数据，10 分钟自动生成可直接上会的调账材料",
        "",
        "✅ 本地知识检索（RAG）",
        "✅ 多步骤任务规划（10-step workflow）",
        "✅ 本地多轮记忆（memory system）",
        "✅ 工具调用（Excel/计算/报告）",
        "✅ 质量校验（7 条规则）",
        "",
        "数据全程不出设备，零云端依赖",
    ]
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        if line.startswith("✅"):
            p.font.size = Pt(20)
            p.font.color.rgb = DARK_GRAY
        elif line.startswith("上传"):
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = DARK_GRAY
        elif line.startswith("数据"):
            p.font.size = Pt(18)
            p.font.color.rgb = AMD_RED
    
    # Slide 4: 10-Step Workflow
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "10 步确定性工作流"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    steps_left = [
        "Step 1: 多格式文件加载",
        "Step 2: 智能噪音识别",
        "Step 3: 基准选择三维论证",
        "Step 4: 扫描日志清洗",
        "Step 5: 盘点结果加载",
    ]
    
    steps_right = [
        "Step 6: 差异金额计算",
        "Step 7: 日志关联三分类",
        "Step 8: 库位状态校验",
        "Step 9: 品类维度分析",
        "Step 10: 交付物生成",
    ]
    
    box_left = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(5))
    tf = box_left.text_frame
    for i, step in enumerate(steps_left):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = step
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(12)
    
    box_right = slide.shapes.add_textbox(Inches(7), Inches(1.5), Inches(6), Inches(5))
    tf = box_right.text_frame
    for i, step in enumerate(steps_right):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = step
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(12)
    
    # Slide 5: Innovation 1 - Noise Detection
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "创新一：噪音自动识别"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    content = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = content.text_frame
    tf.word_wrap = True
    
    lines = [
        "自动识别 3 类噪音：",
        "",
        "1. RF 扫码环境噪声（NOISE）",
        "   → 68 条干扰数据自动过滤",
        "",
        "2. 异常记录",
        "   → 历史库存含 200+ 条异常",
        "",
        "3. 不适用文件",
        "   → 月报粒度不匹配，自动排除",
        "",
        "人工做这一步需要逐行检查，Agent 3 秒完成",
    ]
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
    
    # Slide 6: Innovation 2 - Baseline Selection
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "创新二：基准选择三维论证"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    content = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = content.text_frame
    tf.word_wrap = True
    
    lines = [
        "三维对比自动选对基准：",
        "",
        "1. 编码体系对比",
        "   → 型号编码 vs SKU 编码",
        "",
        "2. 时间窗口对比",
        "   → 快照时间 vs 盘点时间",
        "",
        "3. 数据粒度对比",
        "   → 品类汇总 vs 明细记录",
        "",
        "选错基准 = 整个分析白做",
    ]
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
    
    # Slide 7: Innovation 3 - Log Correlation
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "创新三：日志关联三分类"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    content = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = content.text_frame
    tf.word_wrap = True
    
    lines = [
        "按净变动量精准比对（非简单的有无操作）",
        "",
        "14 个完全可解释",
        "   → 有对应的出入库操作",
        "",
        "21 个部分可解释",
        "   → 有操作但量不足",
        "",
        "124 个无日志关联",
        "   → 需要现场实物复核的重点！",
        "",
        "精准区分，避免遗漏关键问题",
    ]
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
    
    # Slide 8: Demo Results
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "演示结果"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    # Results table
    results = [
        ("系统库存总计", "21,144 件"),
        ("净差异", "-971 件"),
        ("总体差异率", "4.6%"),
        ("净差异金额", "-989,900.97 元"),
        ("日志关联三分类", "14 / 21 / 124"),
        ("质量校验", "7/7 通过"),
    ]
    
    for i, (metric, value) in enumerate(results):
        y = Inches(1.5 + i * 0.8)
        box = slide.shapes.add_textbox(Inches(2), y, Inches(4), Inches(0.6))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = metric
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_GRAY
        p.alignment = PP_ALIGN.RIGHT
        
        box2 = slide.shapes.add_textbox(Inches(7), y, Inches(4), Inches(0.6))
        tf = box2.text_frame
        p = tf.paragraphs[0]
        p.text = value
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = AMD_RED
    
    # Slide 9: AMD GPU
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "为什么选择 AMD ROCm"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    content = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = content.text_frame
    tf.word_wrap = True
    
    lines = [
        "✅ 完全开源 — 无供应商锁定",
        "",
        "✅ 硬件成本更低 — 企业部署友好",
        "",
        "✅ vLLM 原生支持 — 高性能推理",
        "",
        "✅ 数据安全 — 本地部署，数据不出设备",
        "",
        "✅ 零 API 费用 — 一次部署，持续使用",
    ]
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_GRAY
    
    # Slide 10: Value & Contact
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "业务价值"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = AMD_RED
    
    value = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(2))
    tf = value.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "分析耗时：2-3 天 → 10 分钟（提升 99%）"
    p.font.size = Pt(24)
    p.font.color.rgb = DARK_GRAY
    
    contact = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(3))
    tf = contact.text_frame
    tf.word_wrap = True
    
    lines = [
        "GitHub: private repository",
        "PR: AMD-DEV-CONTEST/Radeon-hackathon-2026-07/pull/27",
        "",
        "谢谢！",
    ]
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
    
    # Save
    output_path = "/app/InventoryDoc_Agent_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved: {output_path}")
    return output_path


if __name__ == "__main__":
    create_presentation()
