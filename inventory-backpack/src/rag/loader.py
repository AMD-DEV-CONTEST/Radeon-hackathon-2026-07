"""Document loader - multi-format text extraction with binary detection"""

import os
import re
import logging
from pathlib import Path
from typing import List, Dict
from dataclasses import dataclass

logger = logging.getLogger(__name__)


@dataclass
class Document:
    content: str
    metadata: Dict


class DocumentLoader:
    """Multi-format document loader with binary detection and format-specific parsing"""

    TEXT_EXTENSIONS = {
        '.txt', '.md', '.py', '.js', '.ts', '.java', '.c', '.cpp', '.h',
        '.cs', '.go', '.rs', '.rb', '.php', '.swift', '.kt', '.scala',
        '.html', '.css', '.json', '.xml', '.yaml', '.yml', '.toml',
        '.csv', '.log', '.ini', '.cfg', '.conf', '.sh', '.bat', '.ps1',
        '.sql', '.r', '.m', '.mm', '.pl', '.lua', '.vim', '.ex', '.exs',
    }

    SUPPORTED_EXTENSIONS = {
        '.txt', '.md', '.py', '.js', '.java', '.c', '.cpp', '.h',
        '.pdf', '.docx', '.doc', '.rtf', '.odt',
        '.xlsx', '.xls', '.csv', '.tsv',
        '.pptx',
    }

    BINARY_EXTENSIONS = {
        '.png', '.jpg', '.jpeg', '.gif', '.bmp', '.svg', '.ico',
        '.mp3', '.mp4', '.avi', '.mov', '.mkv', '.wav', '.flac',
        '.zip', '.rar', '.7z', '.tar', '.gz', '.bz2',
        '.exe', '.dll', '.so', '.o', '.a', '.lib', '.class',
        '.pyc', '.pyo', '.elc',
        '.iso', '.bin', '.img', '.dat',
        '.db', '.sqlite', '.sqlite3',
        '.apk', '.ipa', '.jar', '.war', '.ear',
    }

    def __init__(self, data_dir: str = "./data/knowledge_base"):
        self.data_dir = Path(data_dir)

    def load_file(self, file_path: str) -> List[Document]:
        path = Path(file_path)
        if not path.exists():
            logger.warning(f"File not found: {file_path}")
            return []

        ext = path.suffix.lower()

        if ext in self.BINARY_EXTENSIONS:
            logger.info(f"Skipping binary file: {path.name} (type: {ext})")
            return []

        if ext == '.pdf':
            return self._load_pdf(path)
        elif ext in ('.docx', '.doc'):
            return self._load_docx(path)
        elif ext in ('.xlsx', '.xls'):
            return self._load_xlsx(path)
        elif ext in ('.pptx',):
            return self._load_pptx(path)
        elif ext in ('.csv', '.tsv'):
            return self._load_csv(path)
        elif ext in ('.rtf', '.odt'):
            return self._load_text(path)
        else:
            return self._load_text(path)

    def _is_valid_text(self, content: str) -> bool:
        if not content or len(content.strip()) < 5:
            return False
        if self._is_garbage(content):
            return False
        return True

    def _is_garbage(self, content: str) -> bool:
        if not content:
            return True
        non_printable = 0
        high_ratio = 0
        for ch in content:
            code = ord(ch)
            if code < 9 or (code > 13 and code < 32):
                non_printable += 1
            if code == 0xFFFD:
                high_ratio += 1
        if high_ratio > len(content) * 0.1:
            return True
        if non_printable > len(content) * 0.3:
            return True
        replacement_chars = content.count('\ufffd')
        if replacement_chars > len(content) * 0.05:
            return True
        return False

    def _load_text(self, path: Path) -> List[Document]:
        try:
            content = path.read_text(encoding='utf-8')
            if not self._is_valid_text(content):
                logger.info(f"Text file has insufficient content: {path.name}")
                return []
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": path.suffix}
            )]
        except UnicodeDecodeError:
            return self._try_decode(path)
        except Exception as e:
            logger.error(f"Failed to load text file {path}: {e}")
            return []

    def _try_decode(self, path: Path) -> List[Document]:
        raw = path.read_bytes()
        for enc in ['utf-8', 'gbk', 'gb2312', 'latin-1', 'utf-16']:
            try:
                content = raw.decode(enc)
                if self._is_valid_text(content):
                    logger.info(f"Decoded {path.name} with {enc}")
                    return [Document(
                        content=content,
                        metadata={"source": str(path), "filename": path.name, "type": path.suffix}
                    )]
            except (UnicodeDecodeError, LookupError):
                continue
        if self._is_binary_content(raw):
            logger.info(f"Skipping binary content: {path.name}")
            return []
        content = raw.decode('utf-8', errors='replace')
        if self._is_valid_text(content):
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": path.suffix}
            )]
        logger.info(f"File has no usable text content: {path.name}")
        return []

    def _is_binary_content(self, raw: bytes) -> bool:
        if not raw:
            return False
        text_chars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x7F)) | set(range(0x80, 0x100)))
        non_text = raw.translate(None, bytes(text_chars))
        if len(non_text) / max(len(raw), 1) > 0.30:
            return True
        text_test = raw[:8192]
        if b'\x00' in text_test:
            return True
        return False

    def _load_pdf(self, path: Path) -> List[Document]:
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            texts = []
            for page in reader.pages:
                t = page.extract_text()
                if t and t.strip():
                    texts.append(t)
            content = "\n".join(texts)
            if not self._is_valid_text(content):
                logger.info(f"PDF has no extractable text: {path.name}")
                return []
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": ".pdf"}
            )]
        except ImportError:
            logger.warning("pypdf not installed, cannot parse PDF")
            return []
        except Exception as e:
            logger.error(f"PDF load failed {path}: {e}")
            return []

    def _load_docx(self, path: Path) -> List[Document]:
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            content = "\n".join(paragraphs)
            if not self._is_valid_text(content):
                logger.info(f"DOCX has no extractable text: {path.name}")
                return []
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": ".docx"}
            )]
        except ImportError:
            logger.warning("python-docx not installed, cannot parse DOCX")
            return []
        except Exception as e:
            logger.error(f"DOCX load failed {path}: {e}")
            return []

    def _load_xlsx(self, path: Path) -> List[Document]:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(path), data_only=True)
            sheets = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else '' for c in row]
                    row_text = '\t'.join(cells).strip()
                    if row_text:
                        rows.append(row_text)
                if rows:
                    sheets.append(f"=== Sheet: {sheet_name} ===\n" + '\n'.join(rows))
            content = '\n\n'.join(sheets)
            if not self._is_valid_text(content):
                logger.info(f"XLSX has no extractable text: {path.name}")
                return []
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": ".xlsx"}
            )]
        except ImportError:
            logger.warning("openpyxl not installed, cannot parse XLSX")
            return []
        except Exception as e:
            logger.error(f"XLSX load failed {path}: {e}")
            return []

    def _load_pptx(self, path: Path) -> List[Document]:
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text and shape.text.strip():
                        texts.append(shape.text.strip())
                if texts:
                    slides.append(f"=== Slide {i} ===\n" + '\n'.join(texts))
            content = '\n\n'.join(slides)
            if not self._is_valid_text(content):
                logger.info(f"PPTX has no extractable text: {path.name}")
                return []
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": ".pptx"}
            )]
        except ImportError:
            logger.warning("python-pptx not installed, cannot parse PPTX")
            return []
        except Exception as e:
            logger.error(f"PPTX load failed {path}: {e}")
            return []

    def _load_csv(self, path: Path) -> List[Document]:
        try:
            import csv
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            raw = path.read_bytes()
            for enc in encodings:
                try:
                    text = raw.decode(enc)
                    break
                except (UnicodeDecodeError, LookupError):
                    continue
            else:
                text = raw.decode('utf-8', errors='replace')
            rows = []
            reader = csv.reader(text.splitlines())
            for row in reader:
                row_text = '\t'.join(c.strip() for c in row if c.strip())
                if row_text:
                    rows.append(row_text)
            content = '\n'.join(rows)
            if not self._is_valid_text(content):
                logger.info(f"CSV has no extractable text: {path.name}")
                return []
            return [Document(
                content=content,
                metadata={"source": str(path), "filename": path.name, "type": ".csv"}
            )]
        except Exception as e:
            logger.error(f"CSV load failed {path}: {e}")
            return []

    def load_directory(self, dir_path: str = None) -> List[Document]:
        target_dir = Path(dir_path) if dir_path else self.data_dir
        if not target_dir.exists():
            logger.warning(f"Directory not found: {target_dir}")
            return []

        documents = []
        for ext in self.SUPPORTED_EXTENSIONS:
            for file_path in target_dir.rglob(f"*{ext}"):
                docs = self.load_file(str(file_path))
                documents.extend(docs)

        logger.info(f"Loaded {len(documents)} documents from {target_dir}")
        return documents