from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pypdf import PdfWriter

from buildguard.api.app import create_app
from buildguard.config import RagConfig, load_settings
from buildguard.exceptions import ModelUnavailableError
from buildguard.llm.mock_backend import MockLLMBackend
from buildguard.rag.bm25_index import BM25Index
from buildguard.rag.chunker import ChunkingOptions, StructureAwareChunker
from buildguard.rag.citation import make_citation
from buildguard.rag.dense_index import DenseIndex
from buildguard.rag.embedding_backend import (
    MockEmbeddingBackend,
    TransformersEmbeddingBackend,
)
from buildguard.rag.evaluation import load_qa_dataset
from buildguard.rag.fusion import reciprocal_rank_fusion
from buildguard.rag.parsers.docx_parser import DocxParser
from buildguard.rag.parsers.markdown_parser import MarkdownParser
from buildguard.rag.parsers.pdf_parser import PdfParser, split_layout_blocks
from buildguard.rag.parsers.text_parser import TextParser
from buildguard.rag.reranker import MockReranker
from buildguard.rag.schemas import (
    DocumentChunk,
    ParsedDocument,
    ParsedSection,
    QAExample,
    RagSearchInput,
    ScoredChunk,
)
from buildguard.rag.service import RagService
from buildguard.rag.tool import register_rag_tool
from buildguard.tools.registry import ToolRegistry


def rag_config(index_dir: Path) -> RagConfig:
    return RagConfig(
        index_dir=index_dir,
        allowed_source_roots=[index_dir.parent / "knowledge"],
        chunk_min_chars=50,
        chunk_max_chars=120,
        chunk_overlap_chars=10,
        embedding_dimension=64,
        bm25_top_k=20,
        dense_top_k=20,
        fusion_top_k=15,
        final_top_k=5,
        min_top_score=0.05,
        min_keyword_coverage=0.1,
    )


def rag_service(tmp_path: Path) -> RagService:
    config = rag_config(tmp_path / "indexes")
    return RagService(
        config,
        MockEmbeddingBackend(config.embedding_dimension, seed=42),
        MockReranker(),
    )


def write_standard(path: Path) -> None:
    path.write_text(
        "# 第一章 临边防护\n\n"
        "第一条 施工现场临边作业必须设置防护栏杆，栏杆应牢固可靠并设置警示标志。"
        "防护栏杆高度和构造应符合施工安全规范要求。\n\n"
        "第二条 洞口作业必须设置盖板或防护网，盖板应固定并能够承受施工荷载。\n\n"
        "# 第二章 用电安全\n\n"
        "第三条 配电箱应设置漏电保护器，施工用电设备必须可靠接地。",
        encoding="utf-8",
    )


def test_markdown_and_text_parsing_preserve_metadata(tmp_path: Path) -> None:
    markdown = tmp_path / "standard.md"
    write_standard(markdown)
    parsed = MarkdownParser().parse(markdown)
    assert parsed.checksum
    assert any(section.chapter == "第一章" for section in parsed.sections)
    assert any(section.article == "第一条" for section in parsed.sections)
    assert all(section.source_path == markdown.resolve() for section in parsed.sections)

    text = tmp_path / "standard.txt"
    text.write_text("生效日期：2025年1月2日\n\n第一条 安全帽必须正确佩戴。", encoding="utf-8")
    parsed_text = TextParser().parse(text)
    assert parsed_text.sections[0].effective_date is not None
    assert parsed_text.sections[-1].article == "第一条"


def test_docx_parses_paragraphs_and_table_header(tmp_path: Path) -> None:
    path = tmp_path / "standard.docx"
    document = Document()
    document.add_heading("第一章 安全管理", level=1)
    document.add_paragraph("第一条 作业人员必须接受安全教育。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "安全帽"
    table.cell(1, 1).text = "必须佩戴"
    document.save(path)
    parsed = DocxParser().parse(path)
    table_section = next(section for section in parsed.sections if section.is_table)
    assert table_section.table_header == "项目 | 要求"
    assert "安全帽 | 必须佩戴" in table_section.content


def test_empty_pdf_warns_without_ocr(tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as stream:
        writer.write(stream)
    parsed = PdfParser().parse(path)
    assert parsed.sections == []
    assert "OCR was not run" in parsed.warnings[0]
    assert "OCR was not run" in caplog.text


def test_pdf_layout_preserves_multirow_tables_without_reclassifying_prose() -> None:
    blocks = split_layout_blocks(
        [
            "Inspection requirements",
            "Item        Requirement",
            "Helmet      Required",
            "An isolated  aligned phrase is prose",
        ]
    )
    assert blocks == [
        (False, ["Inspection requirements"]),
        (True, ["Item | Requirement", "Helmet | Required"]),
        (False, ["An isolated  aligned phrase is prose"]),
    ]


def test_chunking_never_crosses_articles_and_ids_are_stable(tmp_path: Path) -> None:
    source = tmp_path / "rules.txt"
    source.write_text("rules", encoding="utf-8")
    checksum = "a" * 64
    document = ParsedDocument(
        document_name=source.name,
        source_path=source,
        checksum=checksum,
        sections=[
            ParsedSection(
                document_name=source.name,
                source_path=source,
                article="第一条",
                content="第一条 " + "临边防护必须设置栏杆。" * 20,
                checksum=checksum,
            ),
            ParsedSection(
                document_name=source.name,
                source_path=source,
                article="第二条",
                content="第二条 " + "洞口必须设置盖板。" * 20,
                checksum=checksum,
            ),
        ],
    )
    chunker = StructureAwareChunker(ChunkingOptions(50, 120, 10))
    first = chunker.chunk(document)
    second = chunker.chunk(document)
    assert [chunk.chunk_id for chunk in first] == [chunk.chunk_id for chunk in second]
    assert all(len(chunk.content) <= 120 for chunk in first)
    assert all(not ("第一条" in chunk.content and "第二条" in chunk.content) for chunk in first)


def test_mock_embedding_is_deterministic_and_normalized() -> None:
    backend = MockEmbeddingBackend(64, seed=7)
    first = backend.encode(["临边防护", "施工用电"])
    second = backend.encode(["临边防护", "施工用电"])
    assert np.array_equal(first, second)
    assert np.allclose(np.linalg.norm(first, axis=1), 1)


def test_missing_transformers_model_has_clear_error(tmp_path: Path) -> None:
    with pytest.raises(ModelUnavailableError, match="missing"):
        TransformersEmbeddingBackend(
            tmp_path / "missing-model",
            device="cpu",
            batch_size=1,
        )


def chunks(tmp_path: Path) -> list[DocumentChunk]:
    return [
        DocumentChunk(
            chunk_id="a" * 64,
            document_name="safety.md",
            source_path=tmp_path / "safety.md",
            document_checksum="1" * 64,
            page=1,
            article="第一条",
            content="施工现场临边作业必须设置牢固的防护栏杆",
        ),
        DocumentChunk(
            chunk_id="b" * 64,
            document_name="electric.md",
            source_path=tmp_path / "electric.md",
            document_checksum="2" * 64,
            page=2,
            article="第三条",
            content="施工用电配电箱必须设置漏电保护器并可靠接地",
        ),
    ]


def test_bm25_and_dense_persist_and_search(tmp_path: Path) -> None:
    corpus = chunks(tmp_path)
    bm25 = BM25Index(tmp_path / "index")
    bm25.build(corpus)
    bm25.save()
    loaded_bm25 = BM25Index(tmp_path / "index")
    assert loaded_bm25.load()
    assert loaded_bm25.checksum == bm25.checksum
    assert loaded_bm25.search("临边防护栏杆", 2)[0].chunk.chunk_id == "a" * 64

    embedding = MockEmbeddingBackend(64)
    dense = DenseIndex(tmp_path / "index", embedding)
    dense.build(corpus)
    dense.save()
    loaded_dense = DenseIndex(tmp_path / "index", embedding)
    assert loaded_dense.load()
    results = loaded_dense.search("临边防护栏杆", 2)
    assert results
    assert results[0].chunk.chunk_id == "a" * 64


def test_rrf_formula_and_citation(tmp_path: Path) -> None:
    corpus = chunks(tmp_path)
    bm25 = [
        ScoredChunk(chunk=corpus[0], score=2, rank=1, source="bm25"),
        ScoredChunk(chunk=corpus[1], score=1, rank=2, source="bm25"),
    ]
    dense = [
        ScoredChunk(chunk=corpus[1], score=0.9, rank=1, source="dense"),
        ScoredChunk(chunk=corpus[0], score=0.8, rank=2, source="dense"),
    ]
    fused = reciprocal_rank_fusion([bm25, dense], rrf_k=60, top_k=2)
    assert fused[0].score == pytest.approx(1 / 61 + 1 / 62)
    citation = make_citation(corpus[0])
    assert citation.label == "safety.md · p.1 · 第一条"


def test_service_incremental_duplicate_remove_and_persistence(tmp_path: Path) -> None:
    source = tmp_path / "standard.md"
    write_standard(source)
    service = rag_service(tmp_path)
    first = service.ingest(source, project="A")
    assert first.status == "indexed"
    assert first.chunk_count > 0
    unchanged = service.ingest(source, project="A")
    assert unchanged.status == "unchanged"
    duplicate = tmp_path / "copy.md"
    duplicate.write_bytes(source.read_bytes())
    duplicate_result = service.ingest(duplicate)
    assert duplicate_result.status == "unchanged"
    assert "Duplicate" in duplicate_result.warnings[0]

    restored = rag_service(tmp_path)
    assert restored.status().document_count == 1
    assert restored.status().chunk_count == first.chunk_count
    removed = restored.remove(source)
    assert removed.status == "removed"
    assert restored.status().chunk_count == 0


def test_hybrid_search_empty_and_insufficient_evidence(tmp_path: Path) -> None:
    service = rag_service(tmp_path)
    empty = service.search(RagSearchInput(query="不存在的规范"))
    assert empty.chunks == []
    assert empty.evidence_status == "insufficient_evidence"

    source = tmp_path / "standard.md"
    write_standard(source)
    assert service.ingest(source).status == "indexed"
    found = service.search(RagSearchInput(query="临边防护栏杆"))
    assert found.chunks
    assert found.citations
    rejected = service.search(RagSearchInput(query="GB 99999-2099 塔吊量子规则"))
    assert rejected.evidence_status == "insufficient_evidence"


def test_rag_tool_and_evaluation_metrics(tmp_path: Path) -> None:
    service = rag_service(tmp_path)
    source = tmp_path / "standard.md"
    write_standard(source)
    service.ingest(source)
    registry = ToolRegistry()
    register_rag_tool(registry, service)
    tool = registry.get("search_construction_standards")
    output = tool.handler(RagSearchInput(query="施工用电漏电保护器"))
    assert output.chunks  # type: ignore[attr-defined]
    relevant_id = output.chunks[0].chunk_id  # type: ignore[attr-defined]
    metrics = service.evaluate(
        [QAExample(question="施工用电漏电保护器", relevant_chunk_ids=[relevant_id])]
    )
    assert metrics.evaluated_questions == 1
    assert 0 <= metrics.recall_at_5 <= 1
    assert 0 <= metrics.ndcg_at_10 <= 1


def test_rag_config_validation() -> None:
    with pytest.raises(ValidationError, match="embedding_model_path"):
        RagConfig(embedding_mode="transformers")
    with pytest.raises(ValidationError, match="chunk_overlap_chars"):
        RagConfig(chunk_min_chars=100, chunk_max_chars=200, chunk_overlap_chars=100)


def test_local_qa_dataset_jsonl_format(tmp_path: Path) -> None:
    path = tmp_path / "qa.jsonl"
    path.write_text(
        '{"question":"临边如何防护？","relevant_chunk_ids":["'
        + "a" * 64
        + '"]}\n',
        encoding="utf-8",
    )
    examples = load_qa_dataset(path)
    assert examples[0].question == "临边如何防护？"
    assert examples[0].relevant_chunk_ids == ["a" * 64]


def test_rag_api_ingest_and_search(config_tree: Path, tmp_path: Path) -> None:
    settings = load_settings(config_tree)
    workspace = tmp_path / "workspace"
    knowledge = workspace / "knowledge"
    knowledge.mkdir(parents=True)
    source = knowledge / "standard.md"
    write_standard(source)
    settings.app.workspace_dir = workspace
    settings.app.database_path = workspace / "database" / "test.db"
    settings.rag = rag_config(workspace / "indexes" / "rag").model_copy(
        update={"allowed_source_roots": [knowledge]}
    )
    application = create_app(settings, MockLLMBackend())
    client = TestClient(application)
    ingest = client.post("/rag/ingest", json={"source_path": str(source)})
    assert ingest.status_code == 200
    assert ingest.json()[0]["status"] == "indexed"
    search = client.post("/rag/search", json={"query": "临边防护栏杆", "top_k": 5})
    assert search.status_code == 200
    assert search.json()["chunks"]
